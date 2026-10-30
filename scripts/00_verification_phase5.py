#!/usr/bin/env python3
"""Contre-expertise de la phase 5 : chiffres décisifs du catalogue et des scénarios.

Relit directement les classeurs et PDF primaires (DGAFP, Insee, DEPP, DGFiP, PAP, jaunes,
rapports), recalcule les paramètres dérivés et refait une réforme par une simulation
écrite ici. N'importe ni modele/modele.py ni scripts/11_* ou 12_* : leurs sorties publiées
(registre, catalogue, hypothèses, résultats) sont seulement lues comme valeurs à contrôler.

Écrit data/risques/verification_phase5.csv.
"""
import csv
import gzip
import json
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/risques/verification_phase5.csv"
FIELDS = ["id", "theme", "indicateur", "annee", "valeur_publiee", "valeur_source", "unite", "source_url", "source_localisation", "statut_controle", "note"]
D = ROOT / "data/raw/01/dgafp_ra2026"
U_POSTES = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/ra_2026_postes.zip"
U_EFF = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/ra_2026_effectifs.zip"
U_RET = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/ra_2026_retraite.zip"
U_ES = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/ra_2026_entrants-sortants.zip"
U_REMU = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/ra_2026_remu_niveaux.zip"
U_IP = "https://www.insee.fr/fr/statistiques/fichier/8732435/IP2094.xlsx"
U_DEPP = "https://www.education.gouv.fr/sites/default/files/document/chapitre-3-les-personnels-non-enseignants-520477.xlsx"
U_DGFIP = "https://www.economie.gouv.fr/files/files/directions_services/dgfip/Rapport/2025/ra_2025_cahierstat.pdf"
U_P156 = "https://www.assemblee-nationale.fr/dyn/contenu/visualisation/1087984/file/PAP2026_BG_Gestion_finances_publiques_GA.pdf"
U_P218 = "https://www.budget.gouv.fr/documentation/file-download/31060"
U_JOP = "https://www.budget.gouv.fr/documentation/file-download/34432"
U_DP = "https://presse.economie.gouv.fr/download?n=1076%20-%20DP%20-%20Budget%202027-pdf&id=184723"
U_IA = "https://igas.gouv.fr/sites/igas/files/2026-07/Rapport%20D%C3%A9ploiement%20IA%20(2026).pdf"
U_RAV = "https://presse.economie.gouv.fr/download?n=1900%20-%20Rapport%20de%20Boris%20Ravignon%20sur%20le%20co%C3%BBt%20du%20millefeuille%20administratif-pdf&id=132915"
U_SEN = "https://www.senat.fr/rap/r24-807-1/r24-807-11.pdf"
U_DAE = "https://www.economie.gouv.fr/files/files/directions_services/dae/media-document/Rapport_activite_2025_en_telechargement.pdf"
U_HCAAM = "https://www.strategie.gouv.fr/files/files/Publications/2021%20SP/2022-01-14%20-%20HCAAM%20-%20Quatre%20sc%C3%A9narios%20polaires%20d%27%C3%A9volution%20de%20l%27articulation/rapport_hcaam-_quatre_scenarios_articulation_amo-amc_-_janvier_2022_0.pdf"
U_ISORA = "https://api.imf.org/external/sdmx/2.1/data/ISORA"

rows = []
_wb, _pdf = {}, {}


def put(id, theme, label, year, published, source, unit, url, loc, ok, note="", status="CONFIRME"):
    rows.append(dict(zip(FIELDS, [id, theme, label, year, published, source, unit, url, loc, status if ok else "ECART", note])))


def cell(path, sheet, ref):
    if path not in _wb:
        _wb[path] = openpyxl.load_workbook(path, data_only=True)
    ws = _wb[path][sheet]
    c = ws[ref]
    label = " ".join(str(ws.cell(c.row, j).value) for j in range(1, 4) if ws.cell(c.row, j).value is not None)
    return c.value, label


def flat(text):
    return re.sub(r"[\s  ]", "", text).replace("’", "'").replace("−", "-")


def pdf_has(path, pages, needles):
    if path not in _pdf:
        _pdf[path] = PdfReader(path)
    text = flat("".join(_pdf[path].pages[p - 1].extract_text() or "" for p in pages))
    missing = [n for n in needles if flat(n) not in text]
    return not missing, missing


def text_has(path, needles):
    text = flat(Path(path).read_text(encoding="utf-8", errors="replace"))
    missing = [n for n in needles if flat(n) not in text]
    return not missing, missing


def near(a, b, tol=0.5):
    return abs(float(a) - float(b)) <= tol


def csvdict(path, key):
    with open(ROOT / path, encoding="utf-8") as f:
        return {r[key]: r for r in csv.DictReader(f)}


def main():
    viv = csvdict("data/reformes/registre_viviers.csv", "vivier_id")
    cat = csvdict("data/reformes/catalogue_reformes.csv", "id")
    hyp = json.loads((ROOT / "modele/hypotheses.json").read_text(encoding="utf-8"))
    res_p = csvdict("modele/resultats/reformes_prudent.csv", "id")
    seuils = csvdict("data/reformes/12_seuils.csv", "seuil")

    # --- Bases des viviers
    fpe = D / "ra2026_caracteristiques_postes_donnees_comp_FPE.xlsx"
    corps = {}
    for i, (ref, kw) in enumerate([("B11", "Attachés"), ("B26", "Personnels administratifs et techniques"), ("B37", "Adjoints administratifs")], 1):
        v, lab = cell(fpe, "Figure 2 catégorie", ref)
        corps[ref] = round(v * 1000)
        put(f"V5-0{i}", "Viviers", f"Corps FPE : {lab[:60]}", "fin 2024", {"B11": 79668, "B26": 157099, "B37": 142142}[ref], corps[ref], "PP",
            U_POSTES, f"postes FPE, Figure 2 catégorie!{ref}", kw in lab and corps[ref] == {"B11": 79668, "B26": 157099, "B37": 142142}[ref])
    tot, lab = cell(fpe, "Figure 2 catégorie", "B42")
    put("V5-04", "Viviers", "Champ de la figure des corps FPE", "fin 2024", "FPE tous statuts (supposé)", round(tot * 1000), "PP", U_POSTES,
        "postes FPE, Figure 2 catégorie!B42 (Total)", round(tot * 1000) == 1501736, status="CONFIRME_AVEC_RESERVE",
        note="Le total des corps vaut 1 501 736, soit les seuls fonctionnaires FPE : les contractuels administratifs ne sont pas dans V_FPE_AUTRES. Voir revue B.")
    ok1, m1 = pdf_has(ROOT / "data/cache/02/dgfip_cahier_stat2025.pdf", [46], ["88659"])
    ok2, m2 = pdf_has(ROOT / "data/cache/02/dgfip_cahier_stat2025.pdf", [47], ["5910"])
    put("V5-05", "Viviers", "DGFiP : titulaires et contractuels", "2025", "88 659 + 5 910", "88 659 + 5 910" if ok1 and ok2 else str(m1 + m2), "PP",
        U_DGFIP, "PDF p.46-47", ok1 and ok2, "Les 5 910 contractuels ne figurent pas dans les lignes de corps (fonctionnaires) dont on les soustrait.")
    dpp, lab_d = cell(ROOT / "data/raw/02/depp_ch3_2025_2026.xlsx", "3.1", "D27")
    etp, _ = cell(ROOT / "data/raw/02/depp_ch3_2025_2026.xlsx", "3.1", "L27")
    put("V5-06", "Viviers", "Filière administrative scolaire", "30/11/2025", "52 432 PP ; 50 849 ETP", f"{dpp} PP ; {etp} ETP", "PP / ETP", U_DEPP,
        "tableau 3.1!D27, L27", dpp == 52432 and etp == 50849 and float(viv["V_MEN_ADM"]["etp_base"]) == 50849)
    for vid, (f, sh, ref, pub) in {"V_FPT_ADM": ("ra2026_caracteristiques_postes_donnees_comp_FPT.xlsx", "Figure 1 filiere", "S28", 453224),
                                   "V_FPH_ADM": ("ra2026_caracteristiques_postes_donnees_comp_FPH.xlsx", "Figure 1 filiere", "T25", 121829)}.items():
        v, lab = cell(D / f, sh, ref)
        put(f"V5-0{7 if vid == 'V_FPT_ADM' else 8}", "Viviers", f"{vid} : {lab[:50]}", "fin 2024", pub, round(v * 1000), "PP", U_POSTES, f"{f}!{sh}!{ref}",
            "Administrative" in lab and round(v * 1000) == pub)
    v, lab = cell(D / "RA_2026_effectifs_donnees_comp_FP.xlsx", "Figure 3", "U22")
    put("V5-09", "Viviers", "Organismes de protection sociale", "fin 2024", viv["V_OPS"]["etp_base"], round(v * 1000), "PP", U_EFF, "effectifs FP, Figure 3!U22",
        "protection sociale" in lab and round(v * 1000) == round(float(viv["V_OPS"]["etp_base"])))
    ok, m = pdf_has(ROOT / "data/cache/06/operateurs_plf2027.pdf", [70], ["France Travail", "48232"])
    put("V5-10", "Viviers", "France Travail, plafond d'emplois 2027", "PLF 2027", viv["V_FT"]["etp_base"], "48 232" if ok else str(m), "ETPT", U_JOP, "PDF p.70", ok)
    eq = {k: cell(ROOT / "data/raw/01/insee_IP2094_donnees.xlsx", "Figure 1", r)[0] for k, r in (("FPE", "E11"), ("FPT", "E12"), ("FPH", "E13"))}
    put("V5-11", "Conversion", "Volume de travail EQTP hors contrats aidés", "2024", "2 579,1 / 1 980,4 / 1 222,3 milliers", " / ".join(str(x) for x in eq.values()),
        "milliers d'EQTP", U_IP, "Figure 1!E11:E13", [eq["FPE"], eq["FPT"], eq["FPH"]] == [2579.1, 1980.4, 1222.3])
    pp = {k: cell(D / "ra_2026_vue_emploi_effectifs.xlsx", "Figure 2", r)[0] for k, r in (("FPE", "C6"), ("FPT", "C7"), ("FPH", "C8"))}
    put("V5-12", "Conversion", "Effectifs physiques par versant", "fin 2024", "2 581 870 / 2 017 830 / 1 248 591", " / ".join(str(x) for x in pp.values()),
        "milliers de PP", U_EFF, "vue emploi effectifs, Figure 2!C6:C8", [pp["FPE"], pp["FPT"], pp["FPH"]] == [2581.87, 2017.83, 1248.591])
    bases = {"V_DGFIP": (88659 + 5910) * eq["FPE"] / pp["FPE"],
             "V_FPE_AUTRES": (sum(corps.values()) - 94569 - 52432) * eq["FPE"] / pp["FPE"],
             "V_FPT_ADM": 453224 * eq["FPT"] / pp["FPT"], "V_FPH_ADM": 121829 * eq["FPH"] / pp["FPH"]}
    put("V5-13", "Calcul", "Bases ETP des quatre viviers recalculées", "2024-2025",
        " / ".join(viv[k]["etp_base"] for k in bases), " / ".join(f"{x:.1f}" for x in bases.values()), "ETP", U_POSTES,
        "PP × EQTP/PP du versant", all(near(viv[k]["etp_base"], x, 0.1) for k, x in bases.items()), status="CALCUL_CONFIRME")

    # --- Départs
    pens = {k: cell(D / "ra_2026_vue_flux_de_personnels_retraites.xlsx", "Figure 1", r)[0] for k, r in (("FPE", "E5"), ("FPT", "E10"), ("FPH", "E13"))}
    put("V5-14", "Départs", "Nouvelles pensions de fonctionnaires civils", "2025", "39 146 / 46 897 / 22 753", " / ".join(str(x) for x in pens.values()),
        "milliers", U_RET, "vue retraites, Figure 1!E5, E10, E13", [pens["FPE"], pens["FPT"], pens["FPH"]] == [39.146, 46.897, 22.753])
    fonc = {"FPE": cell(fpe, "Figure 2 catégorie", "B42")[0], "FPT": cell(D / "ra_2026_vue_emploi_caracteristiques_postes.xlsx", "Figure 1", "D11")[0],
            "FPH": cell(D / "ra_2026_vue_emploi_caracteristiques_postes.xlsx", "Figure 1", "D15")[0]}
    put("V5-15", "Départs", "Fonctionnaires par versant", "fin 2024", "1 501 736 / 1 417 086 / 827 014", " / ".join(str(x) for x in fonc.values()),
        "milliers", U_POSTES, "postes FPE B42 ; vue caractéristiques des postes Figure 1!D11, D15", [fonc["FPE"], fonc["FPT"], fonc["FPH"]] == [1501.736, 1417.086, 827.014])
    rates = {k: pens[k] / fonc[k] for k in pens}
    expect = {"V_DGFIP": "FPE", "V_FPE_AUTRES": "FPE", "V_MEN_ADM": "FPE", "V_FPT_ADM": "FPT", "V_FPH_ADM": "FPH", "V_OPS": "FPE", "V_FT": "FPE"}
    put("V5-16", "Calcul", "Taux de départ par vivier", "2025", " / ".join(f"{viv[k]['taux_departs_annuel']}" for k in expect),
        " / ".join(f"{rates[v]:.6f}" for v in expect.values()), "par an", U_RET, "pensions / fonctionnaires",
        all(near(viv[k]["taux_departs_annuel"], rates[v], 1e-6) for k, v in expect.items()), status="CALCUL_CONFIRME",
        note="V_OPS et V_FT reprennent le taux FPE faute de donnée propre (hypothèse affichée).")
    sort, _ = cell(D / "ra_2026_vue_flux_de_personnels_entrees_sorties.xlsx", "Fig. 1", "H22")
    put("V5-17", "Départs", "Sortants fonctionnaires et rapport aux pensions civiles", "2024-2025", "174 655 ; 1,605",
        f"{sort} ; {sort / sum(pens.values()):.4f}", "milliers ; rapport", U_ES, "entrées-sorties Fig. 1!H22 ; somme des pensions", sort == 174.655 and near(sort / sum(pens.values()), 1.605, 0.001),
        status="CALCUL_CONFIRME")

    # --- Coûts
    ok, m = pdf_has(ROOT / "data/cache/01c/pap2026_gestion.pdf", [50], ["55611", "38455", "34425", "64092", "44430", "36025", "31675"])
    put("V5-18", "Coûts", "PAP 156 : coûts HCAS d'entrée et globaux A/B/C, rémunération d'entrée B", "PAP 2026", "55 611 / 38 455 / 34 425 ; 64 092 / 44 430 / 36 025 ; 31 675",
        "présents" if ok else str(m), "€/ETP/an", U_P156, "PDF p.50", ok)
    ok, m = pdf_has(ROOT / "data/cache/01c/pap2026_218.pdf", [24], ["61076", "44531", "43308", "71301", "50235", "44657"])
    put("V5-19", "Coûts", "PAP 218 : coûts HCAS d'entrée et globaux A/B/C", "PAP 2026", "61 076 / 44 531 / 43 308 ; 71 301 / 50 235 / 44 657",
        "présents" if ok else str(m), "€/ETP/an", U_P218, "PDF p.24", ok)
    w = [corps["B11"], corps["B26"], corps["B37"]]
    w = [x / sum(w) for x in w]
    e_fpe = sum(wi * (a + b) / 2 for wi, a, b in zip(w, (55611, 38455, 34425), (61076, 44531, 43308))) / 1000
    m_fpe = sum(wi * (a + b) / 2 for wi, a, b in zip(w, (64092, 44430, 36025), (71301, 50235, 44657))) / 1000
    e_156 = sum(wi * a for wi, a in zip(w, (55611, 38455, 34425))) / 1000
    m_156 = sum(wi * a for wi, a in zip(w, (64092, 44430, 36025))) / 1000
    c = hyp["couts_k_eur"]
    put("V5-20", "Calcul", "Coûts FPE et DGFiP pondérés A/B/C", "2026",
        f"{c['FPE']['cout_entree_hors_pension']} / {c['FPE']['cout_moyen_hors_pension']} ; {c['FPE_DGFIP']['cout_entree_hors_pension']} / {c['FPE_DGFIP']['cout_moyen_hors_pension']}",
        f"{e_fpe:.2f} / {m_fpe:.2f} ; {e_156:.2f} / {m_156:.2f}", "k€", U_P156, "pondération par les corps V5-01 à 03",
        near(c["FPE"]["cout_entree_hors_pension"], e_fpe, 0.01) and near(c["FPE"]["cout_moyen_hors_pension"], m_fpe, 0.01)
        and near(c["FPE_DGFIP"]["cout_entree_hors_pension"], e_156, 0.01) and near(c["FPE_DGFIP"]["cout_moyen_hors_pension"], m_156, 0.01), status="CALCUL_CONFIRME")
    s_fpe, lab1 = cell(D / "ra_2026_remu_niveaux_donnees_comp.xlsx", "Figure 1", "B11")
    s_fpt, lab2 = cell(D / "ra_2026_remu_niveaux_donnees_comp.xlsx", "Figure 1", "B28")
    ratio = s_fpt / s_fpe
    put("V5-21", "Calcul", "Proxy FPT/FPH : rapport des salaires bruts moyens par EQTP", "2024",
        f"0,7645 ; {c['FPT']['cout_entree_hors_pension']} / {c['FPT']['cout_moyen_hors_pension']}",
        f"{ratio:.4f} ; {e_fpe * ratio:.2f} / {m_fpe * ratio:.2f}", "rapport ; k€", U_REMU, "rémunérations, Figure 1!B11 et B28",
        near(ratio, 0.7645, 0.0001) and near(c["FPT"]["cout_entree_hors_pension"], e_fpe * ratio, 0.01) and c["FPH"] == c["FPT"] | {"_source": c["FPH"]["_source"]},
        status="CALCUL_CONFIRME", note="Le proxy FPH reprend la FPT : choix prudent, le salaire moyen FPH inclut les médecins.")
    put("V5-22", "Calcul", "Charges employeur hors pension dans le coût d'entrée B (PAP 156)", "2026", "17,6 %", f"{(38455 - 31675) / 38455:.3%}",
        "% du coût", U_P156, "PDF p.50 : (38 455 − 31 675) / 38 455", near((38455 - 31675) / 38455, 0.1763, 0.0005), status="CALCUL_CONFIRME",
        note="Seule composante mesurée du taux de conservation APU (0,65) ; CSG et retenue pension sont ajoutées par hypothèse.")

    # --- Faits des fiches, relus dans les sources primaires
    with gzip.open(ROOT / "data/raw/04a/isora_selection_2022_2024.csv.gz", "rt", encoding="utf-8") as f:
        iso = {(r["JURISDICTION"], r["INDICATOR"], r["TIME_PERIOD"]): r["OBS_VALUE"] for r in csv.DictReader(f) if r["JURISDICTION"] == "FRA"}
    vals = [float(iso.get(("FRA", k, "2023"), "nan")) for k in ("95230_206", "96980_206", "98810_206", "337_180")]
    put("V5-23", "Fiches R01-R02", "ISORA France : usagers/déclarations, RH, informatique, total ETP fiscaux", "2023", "10 356 / 1 881 / 2 536 / 38 406",
        " / ".join(f"{x:.0f}" for x in vals), "ETP", U_ISORA, "export ISORA, FRA 2023, codes 95230_206, 96980_206, 98810_206, 337_180",
        vals == [10356, 1881, 2536, 38406], note="Correspondance code → fonction reprise de la note 04a, contre-vérifiée au lot 2.")
    t1 = {s: float(cat["R01"][f"taux_{s}"]) for s in ("prudent", "intermediaire", "ambitieux", "radical")}
    calc = {s: 10356 * f / bases["V_DGFIP"] for s, f in zip(t1, (0.05, 0.10, 0.20, 0.30))}
    put("V5-24", "Calcul", "R01 : taux du vivier dérivés de la fonction usagers", "2027-2031", " / ".join(str(x) for x in t1.values()),
        " / ".join(f"{x:.4f}" for x in calc.values()), "part du vivier", U_ISORA, "10 356 × {5,10,20,30 %} / base DGFiP",
        all(near(t1[s], calc[s], 0.00015) for s in t1), status="CALCUL_CONFIRME", note="Arrondis à 4 décimales dans le catalogue.")
    t2 = [float(cat["R02"][f"taux_{s}"]) for s in ("prudent", "intermediaire", "ambitieux", "radical")]
    c2 = [1881 * 0.05 / bases["V_DGFIP"], 4417 * 0.10 / bases["V_DGFIP"], 4417 * 0.15 / bases["V_DGFIP"], 4417 * 0.20 / bases["V_DGFIP"]]
    put("V5-25", "Calcul", "R02 : taux dérivés des fonctions RH et informatique", "2028-2032", " / ".join(str(x) for x in t2), " / ".join(f"{x:.4f}" for x in c2),
        "part du vivier", U_ISORA, "1 881 et 4 417 ETP", all(near(a, b, 0.00015) for a, b in zip(t2, c2)), status="CALCUL_CONFIRME")
    ok, m = text_has(ROOT / "data/raw/06/dgfip_an2026.txt", ["107 202", "91 064", "Effectifs de la DGFiP"])
    put("V5-26", "Fiche R01", "DGFiP : ETPT 2015 et 2025", "2015-2025", "107 202 → 91 064 ; −15,05 %", f"{'présents' if ok else m} ; {(91064 - 107202) / 107202:.2%}",
        "ETPT", "https://www.assemblee-nationale.fr/", "archive data/raw/06/dgfip_an2026 (rapport AN 2026, tableau d'effectifs)", ok, status="CONFIRME")
    ia = ROOT / "data/cache/08/igas_ia2026.pdf"
    ok, m = pdf_has(ia, [618], ["12 875 k€", "508 k€", "France Travail"])
    put("V5-27", "Convention IA", "France Travail : dépenses IA 2025, base de l'analogie 267 €/ETP", "2025", "12 875 k€ ; 508 k€ ; 0,267 / 0,0105 k€ par ETP",
        f"{'présents' if ok else m} ; {12875 / 48232:.4f} / {508 / 48232:.4f}", "k€", U_IA, "PDF p.618, tableau 6", ok and near(12875 / 48232, 0.267, 0.0005),
        note="Diviser une dépense annuelle par l'effectif plafond est une analogie, pas un coût par outil.")
    ok, m = pdf_has(ia, [236], ["226 ETP", "300 ETP", "241 ETP", "diminution des ETP"])
    put("V5-28", "Fiche R26", "France Travail : équivalents ETP et absence de baisse", "sept. 2025", "226 / 300 / 241 ; 767 = 1,59 % du vivier",
        f"{'présents' if ok else m} ; {767 / 48232:.4%}", "ETP", U_IA, "PDF p.236, note 162", ok and near(767 / 48232, 0.0159, 0.0001))
    rav = ROOT / "data/cache/04d/ravignon_rapport2024.pdf"
    ok1, m1 = pdf_has(rav, [44], ["7,4 Md€"])
    ok2, m2 = pdf_has(rav, [96], ["31 040"])
    ok3, m3 = pdf_has(rav, [100], ["102,4", "204,7"])
    put("V5-29", "Fiches R14-R17", "Ravignon : coût de l'enchevêtrement, agents RH, gains RH hypothétiques", "2024", "7,4 Md€ ; 31 040 ; 102,4-204,7 M€",
        "présents" if ok1 and ok2 and ok3 else str(m1 + m2 + m3), "Md€ ; agents ; M€", U_RAV, "PDF p.44, 96, 100", ok1 and ok2 and ok3,
        status="CONFIRME_AVEC_RESERVE", note="Le détail 7,444 Md€ du tableau 10 n'est pas extrait par pypdf ; le texte p.44 donne 7,4 Md€, cohérent.")
    sen = ROOT / "data/cache/04d/senat_agences_rapport2025.pdf"
    ok, m = pdf_has(sen, [231], ["2,7 milliards", "540 millions"])
    put("V5-30", "Fiche R09", "Sénat : fonctions support d'opérateurs et réduction de 20 %", "juillet 2025", "2,7 Md€ ; 540 M€", "présents" if ok else str(m),
        "€", U_SEN, "PDF p.231", ok)
    ok, m = pdf_has(sen, [217], ["20 millions de mètres carrés", "25 %"])
    put("V5-31", "Fiche R29", "Sénat : objectif immobilier sur 20 millions de m²", "2025", "−25 % ; 20 millions de m²", "présents" if ok else str(m),
        "m²", U_SEN, "PDF p.217", ok)
    dae = ROOT / "data/cache/04d/dae_rapport2025.pdf"
    ok1, m1 = pdf_has(dae, [46], ["50 245"])
    ok2, m2 = pdf_has(dae, [48], ["24 364"])
    ok3, m3 = pdf_has(dae, [87], ["412,53"])
    calc28 = [round(24364 * x) for x in (0.0025, 0.005, 0.01, 0.015)]
    pub28 = [int(float(cat["R28"][f"eco_fonct_meur_{s}"])) for s in ("prudent", "intermediaire", "ambitieux", "radical")]
    put("V5-32", "Fiche R28", "Achats de l'État et économies DAE ; assiette de R28", "2025", f"50 245 / 24 364 / 412,53 M€ ; {pub28}",
        f"{'présents' if ok1 and ok2 and ok3 else m1 + m2 + m3} ; {calc28}", "M€", U_DAE, "PDF p.46, 48, 87", ok1 and ok2 and ok3 and calc28 == pub28)
    ok, m = text_has(ROOT / "data/raw/06/rgpp_net_an2012.txt", ["264 millions d'euros en 2010", "808 millions", "373 millions d'euros en 2011"])
    put("V5-33", "Fiche R12", "RGPP : titre 2 brut et net 2010-2011", "2010-2011", "808 → 264 ; 935 → 373 M€", "présents" if ok else str(m), "M€",
        "https://www.assemblee-nationale.fr/", "archive data/raw/06/rgpp_net_an2012", ok)
    ok, m = pdf_has(ROOT / "data/cache/04c/hcaam2022_scenarios.pdf", [79, 80], ["5,4 Md", "22,5 Md"])
    put("V5-34", "Fiche R21", "HCAAM : charges évitées et transfert vers les finances publiques", "janvier 2022", "5,4 ; 22,5 Md€", "présents" if ok else str(m), "Md€",
        U_HCAAM, "PDF p.79-80", ok)
    ok1, m1 = pdf_has(ROOT / "data/cache/06/operateurs_plf2027.pdf", [12], ["423 opérateurs"])
    ok2, m2 = pdf_has(ROOT / "data/cache/06/operateurs_plf2027.pdf", [24], ["72 572"])
    put("V5-35", "Fiche R04", "Opérateurs PLF 2027 et concours publics", "PLF 2027", "423 ; 72,572 Md€", "présents" if ok1 and ok2 else str(m1 + m2), "nombre ; M€",
        U_JOP, "PDF p.12 et 24", ok1 and ok2)
    ok, m = pdf_has(ROOT / "data/cache/10/dp_budget2027.pdf", [58, 59], ["-812", "404", "1076"])
    put("V5-36", "Fiche R06", "PLF 2027 : administration centrale, numérique réinternalisé, total retraité", "PLF 2027", "−812 ; +404 ; −1 076 ETP",
        "présents" if ok else str(m), "ETP", U_DP, "PDF p.58-59", ok, note="Valeurs déjà contrôlées en V10-14 à V10-16 ; relecture de cohérence.")
    emp = D / "RA_2026_effectifs_donnees_comp_FPT_FPH.xlsx"
    got = [round(cell(emp, "Figure 2", r)[0] * 1000) for r in ("R5", "R18", "R24", "R26", "R29")]
    put("V5-37", "Fiches R15, R17", "Employeurs FPT : communes, EPCI, autres intercommunalités, départements, régions", "fin 2024",
        "1 041 290 / 309 577 / 91 825 / 286 800 / 98 916", " / ".join(str(x) for x in got), "PP", U_EFF, "effectifs FPT-FPH, Figure 2!R5, R18, R24, R26, R29",
        got == [1041290, 309577, 91825, 286800, 98916])

    # --- Recalcul indépendant d'une réforme : R17, scénario prudent (seule réforme active de V_FPT_ADM)
    others = [r for r in cat.values() if r["vivier_id"] == "V_FPT_ADM" and r["id"] != "R17" and float(r["taux_prudent"]) > 0]
    base, rate = float(viv["V_FPT_ADM"]["etp_base"]), float(viv["V_FPT_ADM"]["taux_departs_annuel"])
    entry, avg = c["FPT"]["cout_entree_hors_pension"], c["FPT"]["cout_moyen_hors_pension"]
    g, half = hyp["gvt_individuel_reel"], hyp["fraction_economie_premiere_annee"]
    target = float(cat["R17"]["taux_prudent"]) * base
    years = list(range(2027, 2037))
    start, dur = int(cat["R17"]["annee_debut"]), int(cat["R17"]["duree_ans"])
    inv, inv_dur = float(cat["R17"]["invest_meur_prudent"]), int(cat["R17"]["invest_duree_ans"])
    stock, backlog, cohorts, nets = base, 0.0, [], []
    for y in years:
        cap = rate * stock
        backlog += target / dur if start <= y < start + dur else 0
        done = min(backlog, cap)
        backlog -= done
        stock -= done
        cohorts.append((y, done))
        pay = sum((half if y == y0 else 1) * n * min(avg, entry * (1 + g) ** (y - y0)) / 1000 for y0, n in cohorts)
        nets.append(pay - (inv / inv_dur if start <= y < start + inv_dur else 0))
    r17 = res_p["R17"]
    ok = not others and near(r17["etp_cible"], target, 0.01) and near(r17["net_annee_finale_meur"], nets[-1], 0.01) and near(r17["cumul_net_10ans"], sum(nets), 0.01)
    put("V5-38", "Moteur", "R17 prudent refait par une simulation indépendante", "2027-2036",
        f"cible {float(r17['etp_cible']):.1f} ; net 2036 {float(r17['net_annee_finale_meur']):.2f} ; cumul {float(r17['cumul_net_10ans']):.2f}",
        f"cible {target:.1f} ; net 2036 {nets[-1]:.2f} ; cumul {sum(nets):.2f}", "ETP ; M€", "modele/resultats/reformes_prudent.csv",
        "simulation annuelle écrite dans ce script", ok, status="CALCUL_CONFIRME")
    gross = target * entry / 1000
    put("V5-39", "Convention", "R17 prudent : investissement = 2 années d'économie brute au coût d'entrée", "2027-2028", cat["R17"]["invest_meur_prudent"],
        f"{2 * gross:.3f}", "M€", "data/reformes/catalogue_reformes.csv", "2 × cible × 33,67 / 1 000", near(cat["R17"]["invest_meur_prudent"], 2 * gross, 0.001),
        status="CALCUL_CONFIRME")
    # Totaux : la synthèse est-elle la somme des réformes ?
    for k, sc in enumerate(("prudent", "radical"), 40):
        per = csvdict(f"modele/resultats/reformes_{sc}.csv", "id")
        syn = csvdict("modele/resultats/synthese_scenarios.csv", "scenario")[sc]
        s_post = sum(float(r["etp_postes_en_moins"]) for r in per.values())
        s_net = sum(float(r["net_annee_finale_meur"]) for r in per.values())
        put(f"V5-{k}", "Moteur", f"Synthèse {sc} = somme des 30 réformes", "2036", f"{float(syn['postes_en_moins_10ans']):.1f} postes ; {float(syn['eco_nette_2036_meur']):.1f} M€",
            f"{s_post:.1f} postes ; {s_net:.1f} M€", "postes ; M€", "modele/resultats", "somme des lignes reformes_" + sc,
            near(syn["postes_en_moins_10ans"], s_post, 0.01) and near(syn["eco_nette_2036_meur"], s_net, 0.01), status="CALCUL_CONFIRME")
    # Capacités et seuils
    def capa(n, f):
        return sum(float(v["etp_base"]) * (1 - (1 - min(1, float(v["taux_departs_annuel"]) * f)) ** n) for v in viv.values())
    ratio_s = sort / sum(pens.values())
    caps = [capa(5, 1), capa(10, 1), capa(5, ratio_s), capa(10, ratio_s)]
    pub = seuils["500000"]
    pubv = [float(pub[k]) for k in ("capacite_retraites_5ans", "capacite_retraites_10ans", "capacite_toutes_sorties_5ans", "capacite_toutes_sorties_10ans")]
    put("V5-42", "Seuils", "Capacités de départs des viviers administratifs", "2027-2036", " / ".join(f"{x:.0f}" for x in pubv), " / ".join(f"{x:.0f}" for x in caps),
        "postes", "data/reformes/12_seuils.csv", "B × (1 − (1 − r)^n), r × 1 ou × 1,605", all(near(a, b, 1) for a, b in zip(pubv, caps)), status="CALCUL_CONFIRME")
    verdicts = {s: seuils[s]["verdict"] for s in seuils}
    inter, rad, fav = float(pub["catalogue_intermediaire_2036"]), float(pub["catalogue_radical_2036"]), float(pub["radical_favorable_2036"])
    redo = {}
    for s in seuils:
        x = float(s)
        redo[s] = "conditionnel" if rad >= x else "non démontré" if fav >= x or caps[3] >= x else "incompatible"
    put("V5-43", "Seuils", "Verdicts des sept seuils refaits", "2036", json.dumps(verdicts, ensure_ascii=False), json.dumps(redo, ensure_ascii=False), "verdict",
        "data/reformes/12_seuils.csv", "règles de la note 12, section 3", redo == verdicts, status="CALCUL_CONFIRME")
    put("V5-44", "Seuils", "800 000 rapportés aux capacités sur cinq ans", "2027-2031", "5,1 × ; 3,3 ×", f"{800000 / caps[0]:.2f} × ; {800000 / caps[2]:.2f} ×",
        "rapport", "data/reformes/12_seuils.csv", "800 000 / capacité", near(800000 / caps[0], 5.09, 0.01) and near(800000 / caps[2], 3.29, 0.01), status="CALCUL_CONFIRME")
    tot_base = sum(float(v["etp_base"]) for v in viv.values())
    put("V5-45", "Seuils", "Total des viviers administratifs et 35 % de chacun", "2024-2025", "1 151 911 ; 403 169", f"{tot_base:.0f} ; {tot_base * 0.35:.0f}", "ETP",
        "data/reformes/registre_viviers.csv", "somme des bases", near(tot_base, 1151911.3, 0.5) and near(tot_base * 0.35, 403169, 1), status="CALCUL_CONFIRME")
    # Cohérence fiches / catalogue : chaque fiche A–M reprend le taux du catalogue
    bad = []
    for r in cat.values():
        text = (ROOT / r["fiche"]).read_text(encoding="utf-8")
        shown = " / ".join(f"{float(r[f'taux_{s}']) * 100:,.2f} %".replace(".", ",") for s in ("prudent", "intermediaire", "ambitieux", "radical"))
        if shown not in text:
            bad.append(r["id"])
    put("V5-46", "Cohérence", "Taux affichés dans les 30 fiches = catalogue", "2026", "30 fiches", f"{30 - len(bad)} conformes {bad}", "fiches",
        "recherche/fiches", "rubrique F de chaque fiche", not bad, status="CALCUL_CONFIRME")
    with open(ROOT / "data/reformes/catalogue_reformes.csv", encoding="utf-8") as f:
        ids = [r["id"] for r in csv.DictReader(f)]
    put("V5-47", "Cohérence", "Ordres uniques par vivier (non-cumul)", "2026", "aucun doublon",
        str([v for v in viv if len({cat[i]['ordre'] for i in ids if cat[i]['vivier_id'] == v}) != sum(cat[i]['vivier_id'] == v for i in ids)]),
        "vivier", "data/reformes/catalogue_reformes.csv", "colonne ordre", all(len({cat[i]["ordre"] for i in ids if cat[i]["vivier_id"] == v}) == sum(cat[i]["vivier_id"] == v for i in ids) for v in viv),
        status="CALCUL_CONFIRME")

    OUT.parent.mkdir(exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    from collections import Counter
    print(json.dumps({"controles": len(rows), "statuts": Counter(r["statut_controle"] for r in rows),
                      "ecarts": [r["id"] for r in rows if r["statut_controle"] == "ECART"]}, ensure_ascii=False))
    return 1 if any(r["statut_controle"] == "ECART" for r in rows) else 0


if __name__ == "__main__":
    sys.exit(main())
