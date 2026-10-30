#!/usr/bin/env python3
"""Assemble le rapport et la synthèse : HTML autonome, PDF et DOCX, sans pandoc.

Usage : python scripts/build_rapport.py [--pdf]
Entrées : rapport/[0-9][0-9]_*.md (ordre lexicographique), synthese/synthese.md, scripts/rapport.css
Sorties : livrables/rapport_complet.{md,html,docx[,pdf]}, livrables/synthese.{html,docx[,pdf]}
Dépendances : Markdown, python-docx, playwright (Chromium) — voir requirements-recherche.txt.
Les images sont intégrées au HTML ; les SVG sont rendus en PNG pour le DOCX.
"""
import base64
import glob
import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "livrables"
TITLE = ("La France est-elle suradministrée ? Comparaison internationale, audit des effectifs publics "
         "et potentiel de transformation de l'État à l'ère de l'intelligence artificielle")
DATE = "Rapport arrêté au 30 octobre 2026"
MIME = {".svg": "image/svg+xml", ".png": "image/png", ".jpg": "image/jpeg"}


def embed_images(html, base):
    def repl(m):
        src = m.group(1)
        path = (base / src).resolve()
        if src.startswith("http") or not path.exists():
            return m.group(0)
        data = base64.b64encode(path.read_bytes()).decode()
        return f'src="data:{MIME.get(path.suffix, "application/octet-stream")};base64,{data}"'
    return re.sub(r'src="([^"]+)"', repl, html)


def to_html(md_text, title, subtitle, base, toc=True, extra_css=""):
    conv = markdown.Markdown(extensions=["tables", "fenced_code", "toc", "sane_lists"], extension_configs={"toc": {"toc_depth": "1-2"}})
    body = conv.convert(md_text)
    body = re.sub(r'<p><img alt="([^"]*)" src="([^"]+)" ?/?></p>', r'<figure><img alt="\1" src="\2"><figcaption>\1</figcaption></figure>', body)
    body = embed_images(body, base)
    css = (ROOT / "scripts/rapport.css").read_text(encoding="utf-8")
    nav = f'<nav id="TOC"><h2>Sommaire</h2>{conv.toc}</nav>' if toc else ""
    return (f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f"<title>{title}</title><style>{css}{extra_css}</style></head><body>"
            f'<header id="title-block-header"><h1 class="title">{title}</h1><p class="subtitle">{subtitle}</p></header>{nav}{body}</body></html>')


def svg_to_png(svgs, tmp):
    out = {}
    if not svgs:
        return out
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(device_scale_factor=2)
        for svg in svgs:
            text = svg.read_text(encoding="utf-8")
            m = re.search(r'viewBox="0 0 (\d+) (\d+)"', text)
            w, h = (int(m.group(1)), int(m.group(2))) if m else (900, 500)
            page.set_viewport_size({"width": w, "height": h})
            page.set_content(f'<html><body style="margin:0">{text}</body></html>')
            png = tmp / (svg.stem + ".png")
            page.locator("svg").screenshot(path=str(png))
            out[svg] = png
        browser.close()
    return out


def add_runs(par, text):
    """Gras **x**, liens [x](url) réduits à leur texte, code `x`."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part.replace("*", ""))


def to_docx(md_text, path, title, subtitle, base, pngs):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    t = doc.add_heading(title, 0)
    doc.add_paragraph(subtitle).alignment = WD_ALIGN_PARAGRAPH.LEFT
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                r = doc.add_paragraph().add_run(lines[i])
                r.font.name, r.font.size = "Consolas", Pt(8.5)
                i += 1
        elif m := re.match(r"^(#{1,4}) (.*)", line):
            doc.add_heading(re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", m.group(2)).replace("**", ""), min(len(m.group(1)), 4))
        elif m := re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", line):
            img = (base / m.group(2)).resolve()
            img = pngs.get(img, img)
            if img.exists() and img.suffix in (".png", ".jpg"):
                doc.add_picture(str(img), width=Cm(16))
                cap = doc.add_paragraph(m.group(1))
                cap.runs[0].italic = True
        elif line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[i + 1].strip()):
            rows = [line]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
            ncol = len(cells[0])
            table = doc.add_table(rows=len(cells), cols=ncol)
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(cells):
                for c in range(ncol):
                    cell = table.cell(r, c)
                    cell.text = ""
                    add_runs(cell.paragraphs[0], row[c] if c < len(row) else "")
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8)
                        if r == 0:
                            run.bold = True
            continue
        elif m := re.match(r"^\s*([-*]|\d+\.) (.*)", line):
            add_runs(doc.add_paragraph(style="List Bullet" if m.group(1) in "-*" else "List Number"), m.group(2))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    doc.save(path)


def to_pdf(html_path, pdf_path):
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(html_path.resolve().as_uri(), wait_until="networkidle")
        page.pdf(path=str(pdf_path), format="A4", print_background=True,
                 margin={"top": "18mm", "bottom": "18mm", "left": "16mm", "right": "16mm"},
                 display_header_footer=True, header_template="<span></span>",
                 footer_template='<div style="font-size:8px;width:100%;text-align:center;color:#666"><span class="pageNumber"></span> / <span class="totalPages"></span></div>')
        browser.close()


def build(md_text, stem, title, subtitle, base, pdf, toc=True, extra_css=""):
    html = OUT / f"{stem}.html"
    html.write_text(to_html(md_text, title, subtitle, base, toc, extra_css), encoding="utf-8", newline="\n")
    svgs = {(base / s).resolve() for s in re.findall(r"!\[[^\]]*\]\(([^)]+\.svg)\)", md_text)}
    with tempfile.TemporaryDirectory() as tmp:
        pngs = svg_to_png(sorted(svgs), Path(tmp))
        to_docx(md_text, OUT / f"{stem}.docx", title, subtitle, base, pngs)
    if pdf:
        to_pdf(html, OUT / f"{stem}.pdf")
    return html


def main():
    OUT.mkdir(exist_ok=True)
    pdf = "--pdf" in sys.argv
    chapters = sorted(glob.glob(str(ROOT / "rapport/[0-9][0-9]_*.md")))
    full = "\n\n".join(Path(c).read_text(encoding="utf-8").rstrip() for c in chapters) + "\n"
    (OUT / "rapport_complet.md").write_text(full, encoding="utf-8", newline="\n")
    build(full, "rapport_complet", TITLE, DATE + " · Mission d'audit indépendante", ROOT / "rapport", pdf)
    syn = ROOT / "synthese/synthese.md"
    build(syn.read_text(encoding="utf-8"), "synthese", "Synthèse — La France est-elle suradministrée ?", DATE, ROOT / "synthese", pdf, toc=False,
          extra_css="body{font-size:12.5pt;line-height:1.65} table{font-size:10pt} h2{page-break-before:auto;margin-top:1.4em}")
    manifest = {"date": "2026-10-31", "chapitres": [{"fichier": Path(c).relative_to(ROOT).as_posix(), "sha256": hashlib.sha256(Path(c).read_bytes()).hexdigest()} for c in chapters],
                "synthese_sha256": hashlib.sha256(syn.read_bytes()).hexdigest()}
    (OUT / "rapport_complet.manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print("ok", [Path(c).name for c in chapters])


if __name__ == "__main__":
    main()
