"""Audit reproductible SHA HC7 / régimes HF / ESSPROS / comptes DREES.

python -I scripts/04c_sante_protection.py
Sources archivées seulement ; aucune estimation d'économie ni conversion en emplois.
Les blancs restent manquants. Les hiérarchies HF et les sous-totaux français ne sont pas additionnés.
"""
from __future__ import annotations

import csv
import hashlib
import itertools
import json
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/04c"
OUT = ROOT / "data/international"
CHECKS: list[dict] = []


def num(x: str | float | None) -> float | None:
    return None if x is None or x == "" else float(x)


def code(x: str) -> str:
    return x.split(":", 1)[0].strip()


def write(name: str, rows: list[dict]) -> None:
    if not rows:
        raise ValueError(f"Sortie vide : {name}")
    keys = list(dict.fromkeys(k for r in rows for k in r))
    with (OUT / name).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        w.writerows(rows)


def meta(name: str) -> dict:
    return json.loads((RAW / (name + ".meta.json")).read_text(encoding="utf-8"))


def read_sha(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    result = []
    for line, r in enumerate(rows, 2):
        n = {code(k): code(v) for k, v in r.items()}
        n["country_label"] = r["REF_AREA: Reference area"].partition(":")[2].strip()
        n["local_line"] = line
        n["OBS_VALUE"] = float(n["OBS_VALUE"])
        for k, expected in {"FREQ": "A", "MEASURE": "EXP_HEALTH", "FINANCING_SCHEME_REV": "_Z",
                            "MODE_PROVISION": "_T", "PROVIDER": "_T", "FACTOR_PROVISION": "_Z", "ASSET_TYPE": "_Z"}.items():
            assert n[k] == expected, (path, line, k, n[k])
        n["value_base_units"] = n["OBS_VALUE"] * 10 ** int(n["UNIT_MULT"])
        result.append(n)
    return result


def check(id_: str, outcome: bool, detail: str, source: str, location: str) -> None:
    CHECKS.append(dict(id=id_, resultat="CONFIRME" if outcome else "A_EXAMINER", detail=detail,
                       source_url=source, source_localisation=location, annee="2023-2025", statut="ESTIMATION"))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    total_name, scheme_name = "oecd_sha_current_total", "oecd_sha_current_schemes"
    for name in [total_name, scheme_name]:
        assert meta(name)["acces"] == "integral", f"Collecte incomplète : {name}"
    total, schemes = (read_sha(RAW / (name + ".csv")) for name in [total_name, scheme_name])
    total_url, scheme_url = meta(total_name)["source_url"], meta(scheme_name)["source_url"]
    key = lambda r: (r["REF_AREA"], r["TIME_PERIOD"], r["UNIT_MEASURE"], r["FINANCING_SCHEME"], r["FUNCTION"], r["PRICE_BASE"])
    index = {key(r): r for r in total + schemes}
    assert len({key(r) for r in total}) == len(total)
    assert len({key(r) for r in schemes}) == len(schemes)
    # Les extractions recouvrantes doivent être identiques, y compris leur multiplicateur.
    shared = {key(r): r for r in total}
    overlap = 0
    for r in schemes:
        if key(r) in shared:
            s = shared[key(r)]
            assert r["OBS_VALUE"] == s["OBS_VALUE"] and r["UNIT_MULT"] == s["UNIT_MULT"]
            overlap += 1
    check("SHA_RECOUVREMENT", True, f"{overlap} observations communes strictement identiques", scheme_url, "total vs régimes ; clé complète")
    # Comparaison indépendante de la collecte avec les données héritées, pas nouvelle observation statistique.
    legacy = {key(r): r for p in ["oecd_sha_hc7.csv", "oecd_sha_hc7_by_scheme.csv"] for r in read_sha(ROOT / "data/raw/04" / p)}
    changed = [(k, legacy[k]["OBS_VALUE"], r["OBS_VALUE"]) for k, r in index.items() if k in legacy and legacy[k]["OBS_VALUE"] != r["OBS_VALUE"]]
    check("SHA_HERITAGE", not changed, f"{sum(k in legacy for k in index)} observations recoupées ; {len(changed)} différence de valeur", total_url, "CSV primaires collectés le 28/10/2026 vs fichiers hérités")

    def obs(country: str, year: str, unit: str, function: str = "HC7", hf: str = "_T") -> dict | None:
        price = "_Z" if unit.startswith("PT_") else "V"
        return index.get((country, year, unit, hf, function, price))

    def val(country: str, year: str, unit: str, function: str = "HC7", hf: str = "_T") -> float | None:
        r = obs(country, year, unit, function, hf)
        return None if r is None else r["OBS_VALUE"]

    countries = sorted({r["REF_AREA"] for r in total})
    labels = {r["REF_AREA"]: r["country_label"] for r in total}
    comparisons = []
    for year, country in itertools.product(["2023", "2024"], countries):
        row = dict(pays=country, libelle_pays=labels[country], annee=year)
        for function in ["HC7", "HC71", "HC72", "_T"]:
            for unit, suffix in [("PT_EXP_HLTH", "pct_sante"), ("PT_B1GQ", "pct_pib"), ("USD_PPP_PS", "usd_ppa_hab_courants")]:
                row[function + "_" + suffix] = val(country, year, unit, function)
        for hf in ["HF1", "HF11", "HF121", "HF122", "HF2", "HF21"]:
            row[hf + "_HC7_pct_sante_totale"] = val(country, year, "PT_EXP_HLTH", hf=hf)
        row["couverture_HC7"] = "observe" if row["HC7_pct_sante"] is not None else "NON_TROUVE"
        row["flags_source"] = ";".join(sorted({r[k] for r in total if r["REF_AREA"] == country and r["TIME_PERIOD"] == year and r["PRICE_BASE"] in ["V", "_Z"] for k in ["OBS_STATUS", "OBS_STATUS2", "OBS_STATUS3"] if r[k]}))
        row.update(source_url=total_url, source_regimes_url=scheme_url, source_localisation=f"REF_AREA={country};TIME_PERIOD={year};FINANCING_SCHEME=_T;PRICE_BASE=V ou _Z", statut="FAIT", note="Blanc=absence ; HC71/HC72 sont des fonctions, HF11/HF121 des régimes ; ce sont des dépenses, pas des emplois")
        comparisons.append(row)
        for unit, tolerance in [("PT_EXP_HLTH", .00101), ("PT_B1GQ", .00101), ("USD_PPP_PS", .00101)]:
            v = [val(country, year, unit, x) for x in ["HC7", "HC71", "HC72"]]
            if all(x is not None for x in v):
                check(f"HC7_SOMME_{country}_{year}_{unit}", abs(v[0] - v[1] - v[2]) <= tolerance,
                      f"HC7-HC71-HC72={v[0]-v[1]-v[2]:.9f}; tolérance {tolerance}", total_url, f"{country},{year},{unit}")
    write("04c_sha_comparaison.csv", comparisons)

    finance_rows = []
    for r in schemes:
        if r["PRICE_BASE"] not in ["V", "_Z"]:
            continue  # Les volumes Q base 2020 sont présents au brut, jamais mélangés aux courants.
        c, y, u, hf = (r[x] for x in ["REF_AREA", "TIME_PERIOD", "UNIT_MEASURE", "FINANCING_SCHEME"])
        denominator = val(c, y, "XDC", "_T", hf)
        numerator = val(c, y, "XDC", "HC7", hf)
        ratio = None if numerator is None or denominator is None or denominator <= 0 else 100 * numerator / denominator
        finance_rows.append(dict(pays=c, annee=y, fonction=r["FUNCTION"], regime=hf, unite=u,
            prix=r["PRICE_BASE"], valeur=r["OBS_VALUE"], multiplicateur=int(r["UNIT_MULT"]),
            valeur_unites_base=r["value_base_units"], currency_source=r["CURRENCY"],
            ratio_HC7_sur_depense_du_regime_pct=ratio if r["FUNCTION"] == "HC7" else None,
            statut_ratio="ESTIMATION" if ratio is not None and r["FUNCTION"] == "HC7" else "NON_TROUVE",
            flags_source=";".join(r[k] for k in ["OBS_STATUS", "OBS_STATUS2", "OBS_STATUS3"] if r[k]),
            source_url=scheme_url, source_localisation=f"CSV ligne {r['local_line']};{c},{y},{u},{hf},{r['FUNCTION']},{r['PRICE_BASE']}", statut="FAIT",
            note="PT_EXP_HLTH a pour dénominateur toute la santé ; ratio propre au régime recalculé séparément en XDC ; hiérarchie non additive ; XDC n'est pas EUR hors France"))
    write("04c_sha_regimes.csv", finance_rows)

    codelists = {"CL_HEALTH_FUNCTION": ["HC7", "HC71", "HC72", "_T"],
                 "CL_FINANCING_SCHEME": ["_T", "HF1", "HF11", "HF12HF13", "HF121", "HF122", "HF2", "HF21", "HF3"],
                 "CL_UNIT_MEASURE": ["PT_EXP_HLTH", "PT_B1GQ", "USD_PPP_PS", "EUR_PPP_PS", "XDC", "XDC_PS"],
                 "CL_OBS_STATUS": ["E", "P", "B"]}
    nomenclature = []
    for element in ET.parse(RAW / "oecd_sha_structure.xml").getroot().iter():
        if element.tag.rsplit("}", 1)[-1] != "Codelist" or element.get("id") not in codelists:
            continue
        listname = element.get("id")
        for child in element:
            if child.tag.rsplit("}", 1)[-1] != "Code" or child.get("id") not in codelists[listname]:
                continue
            names = [c.text for c in child if c.tag.rsplit("}", 1)[-1] == "Name" and c.get("{http://www.w3.org/XML/1998/namespace}lang") == "en"]
            nomenclature.append(dict(liste=listname, code=child.get("id"), libelle=names[0] if names else "",
                                     annee="nomenclature consultée 2026", source_url=meta("oecd_sha_structure")["source_url"],
                                     source_localisation=f"Codelist={listname};Code={child.get('id')}", statut="FAIT"))
    write("04c_nomenclature.csv", nomenclature)

    # Vérifie le dénominateur des pourcentages HF avec les montants, sans interpréter un blanc comme zéro.
    for c, y, hf in itertools.product(countries, ["2023", "2024"], ["HF1", "HF11", "HF121", "HF122", "HF2", "HF21"]):
        a, d, p = val(c, y, "XDC", "HC7", hf), val(c, y, "XDC", "_T", "_T"), val(c, y, "PT_EXP_HLTH", "HC7", hf)
        if a is not None and d is not None and p is not None:
            calc = 100 * a / d
            check(f"HF_DENOM_{c}_{y}_{hf}", abs(calc-p) <= .00051, f"100*HC7(HF)/santé_totale={calc:.6f}; publié={p}", scheme_url, f"{c},{y},{hf},XDC V et PT_EXP_HLTH")
    # Contrôles d'emboîtement : pas de mise à zéro implicite des composantes absentes.
    for c, y, parent, children in itertools.product(countries, ["2023", "2024"], ["HF1"], [["HF11", "HF121", "HF122"]]):
        numbers = [val(c, y, "XDC", hf=hf) for hf in [parent] + children]
        if all(x is not None for x in numbers):
            delta = numbers[0] - sum(numbers[1:])
            check(f"HF_EMBOITEMENT_{c}_{y}", abs(delta) <= .0021, f"HF1-HF11-HF121-HF122={delta:.9f} millions XDC", scheme_url, f"{c},{y},HC7,XDC,V")

    gaps = []
    for c in countries:
        if c == "FRA":
            continue
        row = dict(pays_comparateur=c, annee=2023)
        fr, peer = val("FRA", "2023", "PT_EXP_HLTH"), val(c, "2023", "PT_EXP_HLTH")
        fr_eur = val("FRA", "2023", "XDC", "_T")
        gap = None if peer is None else fr - peer
        row.update(niveau1_ecart_HC7_points_sante=gap,
                   niveau1_equivalent_arithmetique_france_meur=None if gap is None else gap*fr_eur/100,
                   niveau1_facteur_ratio_france_sur_pays=None if peer is None or peer == 0 else fr/peer,
                   niveau1_ecart_usd_ppa_hab=None if val(c,"2023","USD_PPP_PS") is None else val("FRA","2023","USD_PPP_PS")-val(c,"2023","USD_PPP_PS"))
        for group, children in [("gouvernement_social", ["HF11", "HF121"]), ("assurances_privees", ["HF122", "HF21"])]:
            f = [val("FRA", "2023", "PT_EXP_HLTH", hf=x) for x in children]
            p = [val(c, "2023", "PT_EXP_HLTH", hf=x) for x in children]
            row["niveau2_"+group+"_france_points"] = sum(f) if all(x is not None for x in f) else None
            row["niveau2_"+group+"_pays_points"] = sum(p) if all(x is not None for x in p) else None
            adjusted = sum(f)-sum(p) if all(x is not None for x in f+p) else None
            row["niveau2_ecart_"+group+"_points"] = adjusted
            row["niveau2_equivalent_"+group+"_meur"] = None if adjusted is None else adjusted*fr_eur/100
        row.update(niveau3_ecart_corrige_besoins_qualite=None, statut_niveau3="NON_TROUVE",
                   niveau4_potentiel_economie_publique_meur=None, statut_niveau4="NON_TROUVE",
                   source_url=scheme_url, source_localisation=f"FRA vs {c},2023,HC7,PT_EXP_HLTH et total santé France XDC V", statut="ESTIMATION",
                   note="Montant équivalent = écart de ratio * dépenses françaises ; ce n'est ni une économie ni un intervalle de potentiel. HF11+HF121 est une correction de régimes, pas le secteur institutionnel S13. Blanc si une composante manque.")
        gaps.append(row)
    write("04c_ecarts_et_corrections.csv", gaps)

    # ESSPROS : réponse JSON primaire de 2023 et groupe comparateur hérité, sans pays supplémentaires sélectionnés après résultat.
    d = json.loads((RAW / "eurostat_esspros_2023.json").read_text(encoding="utf-8"))
    cats = [[x for x, _ in sorted(d["dimension"][k]["category"]["index"].items(), key=lambda kv: kv[1])] for k in d["id"]]
    ess = {}
    for n, combo in enumerate(itertools.product(*cats)):
        value = d["value"].get(str(n))
        if value is not None:
            r = dict(zip(d["id"], combo))
            r.update(value=value, flag=d.get("status", {}).get(str(n), ""))
            ess[(r["geo"], r["spdeps"], r["unit"])] = r
    with (ROOT / "data/raw/04/eurostat_spr_exp_type.csv").open(encoding="utf-8-sig", newline="") as f:
        old_ess = list(csv.DictReader(f))
    geos = sorted({r["geo"] for r in old_ess})
    current_ess_url = meta("eurostat_esspros_2023")["source_url"]
    changed, count = [], 0
    for r in old_ess:
        if r["time"] != "2023":
            continue
        new = ess.get((r["geo"], r["spdeps"], r["unit"]))
        assert new is not None
        count += 1
        if new["value"] != float(r["value"]) or new["flag"] != r["flag"]:
            changed.append((r, new))
    check("ESSPROS_HERITAGE", not changed, f"{count} observations primaires 2023 recoupées ; {len(changed)} différence de valeur ou de drapeau", current_ess_url, "JSON-stat vs CSV hérité")
    ess_rows = []
    for g in geos:
        row = dict(pays=g, annee=2023)
        flags = set()
        for spdeps, unit in itertools.product(["ADM", "TOTAL"], ["PC_GDP", "PC_EXP", "PPS_HAB"]):
            s = ess.get((g, spdeps, unit))
            row[spdeps+"_"+unit] = None if s is None else s["value"]
            if s and s["flag"]:
                flags.add(s["flag"])
        row.update(flags_source=";".join(sorted(flags)), source_url=current_ess_url,
                   source_localisation=f"spr_exp_type;freq=A;geo={g};time=2023;spdeps=ADM/TOTAL", statut="FAIT",
                   note="Tous régimes publics et privés ; PPS_HAB=SPA/habitant, pas EUR ; ADM exclut l'administration interne des soins et la politique sociale générale ; chevauche HC7")
        if row["ADM_PPS_HAB"] is not None and row["TOTAL_PPS_HAB"]:
            calc = 100*row["ADM_PPS_HAB"]/row["TOTAL_PPS_HAB"]
            check("ESSPROS_RATIO_"+g, abs(calc-row["ADM_PC_EXP"]) <= .0051, f"100*ADM/TOTAL en SPA={calc:.6f}; PC_EXP={row['ADM_PC_EXP']}", current_ess_url, f"{g},2023,PPS_HAB et PC_EXP")
        ess_rows.append(row)
    write("04c_esspros_comparaison.csv", ess_rows)

    # DREES : lecture des cellules XLSX originales, pas valeurs recopiées de la note.
    wb = load_workbook(RAW / "drees_cns2026_part3.xlsx", data_only=True)
    national_url = meta("drees_cns2026_part3")["source_url"]
    s = wb["F28 - Tab 1"]
    national = []
    for col in "CDEFGH":
        y = int(s[f"{col}4"].value)
        for rown in range(5, 14):
            national.append(dict(indicateur=s[f"B{rown}"].value, valeur=s[f"{col}{rown}"].value,
                                 unite="millions EUR courants", annee=y, champ="Gestion du système de santé, CNS édition 2026 ; sous-totaux imbriqués",
                                 source_url=national_url, source_localisation=f"F28 - Tab 1!{col}{rown};PDF p.179 (page fichier181)", statut="FAIT"))
        assert abs(s[f"{col}13"].value - sum(s[f"{col}{r}"].value for r in [5, 6, 10])) < .00001
        assert abs(s[f"{col}6"].value - sum(s[f"{col}{r}"].value for r in [7, 8, 9])) < .00001
        assert abs(s[f"{col}10"].value - sum(s[f"{col}{r}"].value for r in [11, 12])) < .00001
    api = json.loads((RAW / "drees_cns_gouvernance_records.json").read_text(encoding="utf-8"))["results"]
    api_index = {(r["annee"], r["financeur"]): r for r in api}
    for y, col in [("2023", "F"), ("2024", "G"), ("2025", "H")]:
        for financeur, rown in [("f000", 13), ("f110", 5), ("f120", 10), ("f210", 6)]:
            v = api_index[(y, financeur)]["montant"]/1e6
            check(f"DREES_XLSX_API_{y}_{financeur}", abs(v-s[f"{col}{rown}"].value) < .00001,
                  f"API en EUR / 1e6={v}; XLSX en millions EUR={s[f'{col}{rown}'].value}", national_url,
                  f"F28 - Tab 1!{col}{rown};API poste=q30000,financeur={financeur},annee={y}")
        if y != "2025":
            difference = s[f"{col}13"].value-val("FRA", y, "XDC")
            national.append(dict(indicateur="Ecart de millésime DREES2026 moins SHA collecté", valeur=difference, unite="millions EUR courants", annee=y,
                                 champ="Même année, publications différentes ; cause détaillée non identifiée pour HC7", source_url=national_url,
                                 source_localisation=f"F28 - Tab 1!{col}13 moins SHA FRA HC7 XDC V", statut="ESTIMATION"))
    for rown in [5, 6, 10, 13]:
        national.append(dict(indicateur="Variation 2025/2024 : "+s[f"B{rown}"].value, valeur=s[f"I{rown}"].value,
                             unite="% nominal", annee=2025, champ="Millésime unique CNS2026", source_url=national_url,
                             source_localisation=f"F28 - Tab 1!I{rown}", statut="FAIT"))
    sg = wb["F28 - Graph 2"]
    for col in "CDEFGH":
        for rown in range(5, 9):
            national.append(dict(indicateur=f"{sg[f'{col}4'].value} : {sg[f'B{rown}'].value}", valeur=sg[f"{col}{rown}"].value,
                                 unite="% des cotisations hors taxes", annee=2025, champ="OC frais de soins ; subventions d'exploitation non retirées",
                                 source_url=national_url, source_localisation=f"F28 - Graph 2!{col}{rown};PDF p.181 (page fichier183)", statut="FAIT"))
    woc = load_workbook(BytesIO((RAW / "drees_oc2025.zip").read_bytes()), data_only=True)
    soc = woc["A2.1.4"]
    ocurl = meta("drees_oc2025")["source_url"]
    for year, col in [(2023, "E"), (2024, "H")]:
        for rown in [6, 9, 10, 11, 12, 13]:
            national.append(dict(indicateur="OC compte technique : "+soc[f"B{rown}"].value, valeur=soc[f"{col}{rown}"].value,
                                 unite="millions EUR courants", annee=year, champ="Frais de soins ; cotisations hors taxes ; rapport OC2025, arrêté spécifique",
                                 source_url=ocurl, source_localisation=f"A2.1.4!{col}{rown}", statut="FAIT"))
        fee = sum(soc[f"{col}{rown}"].value for rown in [10, 11, 12])
        national.append(dict(indicateur="OC frais gestion somme 3 postes", valeur=fee, unite="millions EUR courants", annee=year,
                             champ="Sinistres+acquisition+administration ; autres charges H13/E13 non ajoutées", source_url=ocurl,
                             source_localisation=f"A2.1.4!{col}10+{col}11+{col}12", statut="ESTIMATION"))
    for sheet, cells in [("4.2", ["P5", "P6", "P7"]), ("4.3", ["D16", "E16", "F16", "G16"])]:
        for cell in cells:
            national.append(dict(indicateur="OC taux gestion : "+str(woc[sheet][f"B{woc[sheet][cell].row}"].value), valeur=woc[sheet][cell].value,
                                 unite="% des cotisations hors taxes", annee=2024, champ="Rapport OC2025 ; ligne/colonne détaillée dans le XLSX",
                                 source_url=ocurl, source_localisation=f"{sheet}!{cell}", statut="FAIT"))
    write("04c_france_couts.csv", national)
    check("DREES_IDENTITES", True, "Total=SS+OC+APUC ; OC=3 familles ; APUC=Etat+opérateurs, sur les 6 années publiées", national_url, "F28 - Tab 1!C5:H13")
    write("04c_controles.csv", CHECKS)

    reforms = [
        ("04c_R01", "Simplifier le circuit des droits et remboursements AMO-AMC", "Caisses, OC, professionnels de santé",
         "Interfaces et normes, reprise des données, cybersécurité, formation, double fonctionnement",
         "Mesure de volumes et de temps par tâche ; droits, identifiants, responsabilités et contrôle des rejets partagés",
         "Minutes et coût complet par dossier;rejets et reprises;délais de paiement;créances;incidents de droits",
         "Gains chez payeurs et prestataires à distinguer ; une partie reste hors HC7 ; déduire tâches reprises",
         "https://www.strategie-plan.gouv.fr/files/files/Publications/2021%20SP/2022-01-14%20-%20HCAAM%20-%20Quatre%20sc%C3%A9narios%20polaires%20d%27%C3%A9volution%20de%20l%27articulation/document_travail_hcaam_-_complementaire_sante_et_prevoyance_-_janvier_2021.pdf", "HCAAM2021 pp.39 et41"),
        ("04c_R02", "Réduire la complexité contractuelle et les doublons de commercialisation", "OC, employeurs et assurés",
         "Refonte contrats et appels d'offres, migration, information et accompagnement",
         "Garanties comparables et suivi des personnes fragiles ; examiner concentration, sélection et transfert des primes",
         "Frais par personne et par cotisation;couverture retraités et précaires;primes à panier constant;réclamations;résiliations",
         "Frais privés observés ne sont pas une masse salariale APU ; économies de cotisations employeur public possibles mais non quantifiées",
         ocurl, "RapportOC2025,4.2 et4.3 ; CNS2026 F28 - Graph2"),
        ("04c_R03", "Clarifier le partage AMO-AMC avec simulation redistributive récente", "État, caisses, OC, employeurs, ménages",
         "Reprises de tâches et de personnels, accompagnement social, migration SI, cotisations et fiscalité, prestations supplémentaires",
         "Micro-simulation actualisée, paniers et financement explicites ; droits conservés ; trajectoire de transition",
         "Reste à charge par âge et décile;non-recours;couverture;coût net ménages employeurs APU;emploi transféré",
         "5,4 Md€ du HCAAM est une estimation historique conditionnelle, non un paramètre2027 ; distinguer gain privé et charge publique",
         "https://www.strategie-plan.gouv.fr/files/files/Publications/2021%20SP/2022-01-14%20-%20HCAAM%20-%20Quatre%20sc%C3%A9narios%20polaires%20d%27%C3%A9volution%20de%20l%27articulation/rapport_hcaam-_quatre_scenarios_articulation_amo-amc_-_janvier_2022_0.pdf", "HCAAM2022 pp.73-80, scénario3"),
        ("04c_R04", "Améliorer la gestion des indus et le paiement à bon droit", "Caisses du régime général et assurés",
         "Outils intégrés, conservation des données légales, formation, contrôle et recours",
         "Suivi par cohorte et protection des droits ; mesurer le rendement additionnel net des coûts, pas le recouvrement déjà réalisé",
         "Erreurs de première liquidation;indus et rappels;recouvrement à24mois;coût unitaire;annulations;succès des recours;précarité",
         "Gain de recouvrement distinct d'économie HC7 ; ne pas compter simultanément rendement et suppression des moyens nécessaires",
         "https://www.ccomptes.fr/sites/default/files/2025-05/20250526-RALFSS-2025-Gestion-indus-par-caisses-SS-du-regime-general.pdf", "RALFSS2025 pp.295 et310-313, passages primaires lus via web"),
        ("04c_R05", "Mutualiser des fonctions support réellement communes de pilotage santé", "Ministère, opérateurs, caisses",
         "Diagnostic mission/processus, intégration SI, mobilité, double fonctionnement",
         "Cartographie disjointe des missions ; maintien de régulation, contrôle et accueil ; examen de l'organisation destinataire",
         "Coût par processus;disponibilité SI;délai de réponse;qualité de contrôle;erreurs et ruptures de service",
         "Budgets complets des ARS/opérateurs non assimilables à des économies ; déduire transferts et nouveaux contrats",
         national_url, "CNS2026 fiche28 p.179 ; F28 - Tab1!H10:H12")
    ]
    reform_rows = []
    for id_, label, actors, costs, conditions, kpis, scope, url, location in reforms:
        reform_rows.append(dict(id=id_, intitule=label, acteurs=actors, annee="à instruire pour2027-2036",
                               cout_transition_meur=None, economie_nette_publique_meur=None, etp_evitables=None,
                               statut_chiffrage="NON_TROUVE", couts_a_instruire=costs, conditions=conditions,
                               indicateurs_service=kpis, perimetre_et_double_compte=scope,
                               source_url=url, source_localisation=location, statut="HYPOTHESE"))
    write("04c_pistes_reformes.csv", reform_rows)

    manifest = []
    for folder in [RAW, ROOT / "data/cache/04c"]:
        for path in sorted(folder.glob("*.meta.json")):
            m = json.loads(path.read_text(encoding="utf-8"))
            manifest.append(dict(fichier_manifeste=str(path.relative_to(ROOT)), **m))
    (RAW / "manifestes_collecte.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    # Catalogue raisonné : les empreintes des lectures web portent sur leur export JSON, jamais sur un PDF non téléchargé.
    catalog = json.loads((RAW / "catalogue_sources.json").read_text(encoding="utf-8"))
    registry = []
    for entry in catalog:
        mpath = ROOT / entry["meta_path"]
        m = json.loads(mpath.read_text(encoding="utf-8")) if mpath.exists() else {}
        row = {k: entry.get(k, "") for k in ["id", "titre", "institution", "date_publication"]}
        row.update(url=entry.get("url", m.get("source_url", "")), date_consultation="2026-10-28",
                   acces=entry.get("acces", m.get("acces", "inaccessible")),
                   sha256_16=m.get("sha256", "")[:16], notes=entry.get("notes", ""))
        if entry.get("web_export"):
            row["sha256_16"] = hashlib.sha256((ROOT / entry["web_export"]).read_bytes()).hexdigest()[:16]
        registry.append(row)
    with (ROOT / "sources/registre_04c.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["id", "titre", "institution", "date_publication", "url", "date_consultation", "acces", "sha256_16", "notes"])
        w.writeheader()
        w.writerows(registry)
    failures = [r for r in CHECKS if r["resultat"] != "CONFIRME"]
    print(json.dumps({"lignes_sha_total": len(total), "lignes_sha_regimes": len(schemes), "pays": len(countries),
                      "comparaisons": len(comparisons), "esspros_pays": len(ess_rows), "indicateurs_france": len(national),
                      "controles": len(CHECKS), "a_examiner": len(failures)}, ensure_ascii=False))
    for r in failures:
        print(r["id"], r["detail"])


if __name__ == "__main__":
    main()
