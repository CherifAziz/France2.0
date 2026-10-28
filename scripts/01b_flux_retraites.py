#!/usr/bin/env python3
"""Extrait les flux, pensions et âges du RA DGAFP 2026, sans télécharger.

Exécution depuis toute localisation : python -I scripts/01b_flux_retraites.py
Dépendance : openpyxl. Données : personnes physiques ou pensions, jamais ETP.
Les extrapolations sont des tests mécaniques et non des prévisions démographiques.
"""
from __future__ import annotations

import csv
import hashlib
import json
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/01/dgafp_ra2026"
OUT = ROOT / "data/france"
DATE = "2026-10-28"
BASE = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/"
URL_FLUX = BASE + "ra_2026_entrants-sortants.zip"
URL_RETRAITES = BASE + "ra_2026_retraite.zip"
URL_AGE = "https://www.fonction-publique.gouv.fr/toutes-les-publications/les-agents-de-50-ans-ou-plus-dans-la-fonction-publique-en-2024"
FILES = {
    "flux": "ra_2026_vue_flux_de_personnels_entrees_sorties.xlsx",
    "comp": "ra_2026_entrees_sorties_donnees_comp.xlsx",
    "retraites": "ra_2026_vue_flux_de_personnels_retraites.xlsx",
    "comp_retraites": "RA2026_Retraites_donnees_comp_flux.xlsx",
    "age": "ra_2026_50plus.xlsx",
}
SOURCE_URLS = {"flux": URL_FLUX, "comp": URL_FLUX,
               "retraites": URL_RETRAITES, "comp_retraites": URL_RETRAITES, "age": URL_AGE}
SRE_URL = "https://retraitesdeletat.gouv.fr/statistiques/page-3fp.html"
SRE_CACHE = ROOT / "data/cache/01b/sre_2025_trois_fonctions_publiques.html"


class TableReader(HTMLParser):
    """Lit les tables HTML sauvegardées, sans JavaScript ni accès réseau."""
    def __init__(self):
        super().__init__()
        self.depth = 0
        self.tables = []
        self.rows = []
        self.row = []
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.depth += 1
            if self.depth == 1:
                self.rows = []
        if self.depth == 1:
            if tag == "tr":
                self.row = []
            if tag in ("td", "th"):
                self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if self.depth == 1:
            if tag in ("td", "th") and self.cell is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.cell = None
            if tag == "tr" and self.row:
                self.rows.append(self.row)
        if tag == "table":
            if self.depth == 1:
                self.tables.append(self.rows)
            self.depth -= 1


def read_sre_cache():
    if not SRE_CACHE.exists():
        return [], {}
    parser = TableReader()
    parser.feed(SRE_CACHE.read_text(encoding="utf-8", errors="replace"))
    records = []
    for index, title, years in [(0, "Effectif nouvelles pensions / motif de départ", range(2021, 2026)),
                                 (3, "Effectif nouvelles pensions / domaine administratif", range(2023, 2026))]:
        table = parser.tables[index]
        assert table[0][1:] == [str(year) for year in years], "Structure du tableau SRE modifiée"
        group = ""
        for row in table[2:]:
            if row[0]:
                group = row[0]
            subcategory = row[1] or "Tous motifs"
            for offset, year in enumerate(years):
                records.append({"indicateur": "nouvelles_pensions_droit_direct", "groupe": group,
                                "sous_categorie": subcategory, "annee": year,
                                "valeur": int(row[2 + 2 * offset].replace(" ", "")),
                                "unite": "pensions", "statut": "FAIT", "source_url": SRE_URL,
                                "source_localisation": f"{title} / {group} / {subcategory} / {year}",
                                "note": "FPE civile comprend La Poste et Orange ; FPH et FPT : CNRACL. Même source administrative que DGAFP, publication primaire distincte."})
    def value(group, sub, year, section):
        return next(rec["valeur"] for rec in records if rec["groupe"] == group and rec["sous_categorie"] == sub
                    and rec["annee"] == year and section in rec["source_localisation"])
    reconciliation = {
        "SRE_FPE_civils_2025_y_compris_Poste_Orange": value("FPE civils", "Tous motifs", 2025, "motif"),
        "SRE_La_Poste_2025": value("FPE civils", "La Poste", 2025, "domaine"),
        "SRE_Orange_2025": value("FPE civils", "Orange", 2025, "domaine"),
        "SRE_FPT_2025": value("FPT", "Tous motifs", 2025, "motif"),
        "SRE_FPH_2025": value("FPH", "Tous motifs", 2025, "motif"),
        "SRE_militaires_2025": value("FPE militaires", "Tous motifs", 2025, "motif"),
    }
    assert reconciliation["SRE_FPE_civils_2025_y_compris_Poste_Orange"] - reconciliation["SRE_La_Poste_2025"] - reconciliation["SRE_Orange_2025"] == 39146
    assert reconciliation["SRE_FPT_2025"] == 46897
    assert reconciliation["SRE_FPH_2025"] == 22753
    assert reconciliation["SRE_militaires_2025"] == 10691
    return records, reconciliation


def write_csv(name, rows):
    path = OUT / name
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def thousands(value):
    if not isinstance(value, (float, int)):
        raise ValueError(f"Nombre attendu, reçu {value!r}")
    return int(round(value * 1000))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    books = {key: load_workbook(RAW / name, read_only=True, data_only=True)
             for key, name in FILES.items()}

    def record(key, sheet, cell, indicator, group, dimension, year,
               unit="personnes_physiques", multiplier=1000, note=""):
        value = books[key][sheet][cell].value
        return {
            "indicateur": indicator, "groupe": group, "dimension": dimension,
            "valeur": thousands(value) if multiplier == 1000 else value,
            "unite": unit, "annee": year, "statut": "FAIT",
            "source_url": SOURCE_URLS[key],
            "source_localisation": f"{FILES[key]} / {sheet}!{cell}",
            "note": note,
        }

    # Figure 2 : cellules en milliers, séries 2016-2024.
    history = []
    rows = {5: ("FP_civile", "total"), 6: ("FPE", "versant"),
            7: ("FPT", "versant"), 8: ("FPH", "versant"),
            9: ("Fonctionnaires", "statut"), 10: ("Contractuels", "statut"),
            11: ("Autres", "statut"), 12: ("Femmes", "sexe"), 13: ("Hommes", "sexe")}
    for year in range(2016, 2025):
        col = 2 + 2 * (year - 2016)
        for row, (group, dim) in rows.items():
            for offset, indicator in [(0, "entrants"), (1, "sortants")]:
                cell = books["comp"]["Figure 2"].cell(row, col + offset).coordinate
                history.append(record("comp", "Figure 2", cell, indicator, group, dim, year,
                                      note="Civils, hors contrats aidés, France hors Mayotte ; états au 31/12. Ruptures DSN ; ne pas sommer les dimensions."))
    write_csv("01b_flux_annuels.csv", history)

    # Figure 2 de la vue : croisement complet versant/statut en 2024.
    cross = []
    for branch, start in [("FPE", 6), ("FPT", 11), ("FPH", 16)]:
        for offset, status in enumerate(["Fonctionnaires", "Contractuels", "Autres", "Ensemble"]):
            for col, indicator, unit, mult in [("D", "entrants", "personnes_physiques", 1000),
                                                ("H", "sortants", "personnes_physiques", 1000),
                                                ("F", "taux_entree", "pourcent", 1),
                                                ("J", "taux_sortie", "pourcent", 1)]:
                cross.append(record("flux", "Fig. 2", f"{col}{start + offset}", indicator,
                                    f"{branch}/{status}", "versant_statut", 2024, unit, mult,
                                    "Mouvements interversants et interstatuts exclus des entrées/sorties nettes de FP civile."))
    write_csv("01b_flux_2024_versant_statut.csv", cross)

    movements = []
    balances = {}
    for branch, row in [("FPE", 11), ("FPT", 16), ("FPH", 21)]:
        values = {}
        for col, indicator, year in [("D", "stock_31_decembre", 2023),
                                     ("K", "stock_31_decembre", 2024),
                                     ("E", "entrants_externes", 2024),
                                     ("I", "sortants_externes", 2024),
                                     ("H", "entrants_autres_versants", 2024),
                                     ("J", "sortants_autres_versants", 2024)]:
            rec = record("flux", "Encadré 1", f"{col}{row}", indicator, branch, "versant", year,
                         note="Les mouvements entre versants sont internes à la FP ; ils ne réduisent pas l'emploi public national.")
            movements.append(rec)
            values[col] = rec["valeur"]
        balances[branch] = values["K"] - (values["D"] + values["E"] - values["I"] + values["H"] - values["J"])
    write_csv("01b_mobilites_2024.csv", movements)

    pensions = []
    groups = {6: "FPE_fonctionnaires_civils_SRE", 7: "FPE_militaires_SRE", 8: "FPE_ouvriers_FSPOEIE",
              9: "FPT_fonctionnaires_CNRACL", 10: "FPH_fonctionnaires_CNRACL", 11: "FP_Ircantec"}
    for col in range(3, 29):
        year = int(books["retraites"]["Figure 2"].cell(5, col).value)
        for row, group in groups.items():
            cell = books["retraites"]["Figure 2"].cell(row, col)
            if cell.value is not None:
                pensions.append(record("retraites", "Figure 2", cell.coordinate, "nouvelles_pensions_droit_direct",
                                       group, "regime_statut", year, "pensions",
                                       note="Hors La Poste/Orange pour SRE ; Ircantec : activité publique moins d'un an avant liquidation, disponible depuis 2008. Ne pas assimiler pension à un ETP vacant."))
    # Décomposition 2025, servant aussi au contrôle indépendant de la série.
    latest = []
    for row, group in [(5, "FPE_fonctionnaires_civils_SRE"), (6, "FPE_militaires_SRE"),
                       (7, "FPE_ouvriers_FSPOEIE"), (8, "FPE_contractuels_Ircantec"),
                       (9, "FPE_autres_Ircantec"), (10, "FPT_fonctionnaires_CNRACL"),
                       (11, "FPT_contractuels_Ircantec"), (12, "FPT_autres_Ircantec"),
                       (13, "FPH_fonctionnaires_CNRACL"), (14, "FPH_contractuels_Ircantec"),
                       (15, "FPH_autres_Ircantec"), (16, "FP_hors_militaires"), (17, "FP_total")]:
        latest.append(record("retraites", "Figure 1", f"E{row}", "nouvelles_pensions_droit_direct",
                             group, "total" if row > 15 else "regime_statut", 2025, "pensions",
                             note="Totaux emboîtés : ne pas additionner FP_total et ses composantes. Périmètre des régimes, non ETP libérés."))
    write_csv("01b_retraites_annuelles.csv", pensions)
    write_csv("01b_retraites_2025.csv", latest)

    ages = []
    for sheet, indicator in [("Figure 1", "agents_50_ans_ou_plus"),
                             ("Figure 4", "fonctionnaires_50_ans_ou_plus"),
                             ("Figure 7", "fonctionnaires_60_ans_ou_plus")]:
        for row, branch in [(6, "FPE"), (7, "FPT"), (8, "FPH"), (9, "FP_total")]:
            ages.append(record("age", sheet, f"C{row}", indicator, branch, "versant", 2024,
                               note="Stock d'âge au 31/12 ; ne constitue pas un flux de retraite 2027-2031."))
            ages.append(record("age", sheet, f"D{row}", indicator + "_part", branch, "versant", 2024,
                               "pourcent", 1, "Part dans le champ de la figure ; ne pas utiliser comme probabilité de départ."))
    write_csv("01b_ages_2024.csv", ages)

    # Tests mécaniques : les années 2027-2031 n'ont aucune observation dans ces fichiers.
    # Population de départ supposée égale à fin 2024 (pas d'estimation cachée de 2025-2026).
    projections = []
    metrics = {}
    for branch, row, flowrow in [("FPE", 11, 9), ("FPT", 16, 14), ("FPH", 21, 19)]:
        base = thousands(books["flux"]["Encadré 1"][f"K{row}"].value)
        flow = thousands(books["flux"]["Fig. 2"][f"H{flowrow}"].value)
        rate = flow / base  # Rapport au stock final, pas au stock moyen de la source.
        remain = float(base)
        cum = 0.0
        for year in range(2027, 2032):
            departure = remain * rate
            remain -= departure
            cum += departure
            projections.append({"annee": year, "groupe": branch, "scenario": "zero_recrutement_taux_sortie_constant",
                                "stock_debut_suppose_pp": round(remain + departure, 6),
                                "flux_sortie_suppose_pp": round(departure, 6), "sorties_cumulees_pp": round(cum, 6),
                                "taux_sortie_applique": rate, "unite": "personnes_physiques", "statut": "HYPOTHÈSE",
                                "source_url": URL_FLUX,
                                "source_localisation": f"{FILES['flux']} / Encadré 1!K{row} et Fig. 2!H{flowrow}",
                                "note": "Test maximal tous métiers civils ; stock début 2027 supposé égal fin 2024 ; aucun recrutement ; taux flux 2024/stock fin 2024 ; ne prouve aucun départ définitif administratif ni économie soutenable."})
        metrics[branch] = {"base": base, "flux": flow, "taux_rapport_stock_fin": rate, "sorties_5ans_sans_recrutement": cum}
    write_csv("01b_projection_2027_2031.csv", projections)

    # Comparaison de bases, sans attribuer leur choix à un auteur non vérifié.
    last_exits = next(r["valeur"] for r in history if r["groupe"] == "FP_civile" and r["annee"] == 2024 and r["indicateur"] == "sortants")
    last_titulars = next(r["valeur"] for r in history if r["groupe"] == "Fonctionnaires" and r["annee"] == 2024 and r["indicateur"] == "sortants")
    historical_five = sum(r["valeur"] for r in history if r["groupe"] == "FP_civile" and 2020 <= r["annee"] <= 2024 and r["indicateur"] == "sortants")
    pensions_total = sum(r["valeur"] for r in pensions if r["annee"] == 2025)
    pensions_civil = pensions_total - thousands(books["retraites"]["Figure 1"]["E6"].value)
    bases = [
        ("cumul_historique_sorties_civiles_2020_2024", historical_five, "ESTIMATION", "Somme des cinq flux observés, non personnes uniques ; ruptures DSN."),
        ("cinq_fois_sorties_civiles_2024", 5 * last_exits, "HYPOTHÈSE", "Flux annuel constant, sans correction du stock décroissant ; tous métiers et tous statuts civils."),
        ("cinq_fois_sorties_fonctionnaires_2024", 5 * last_titulars, "HYPOTHÈSE", "Tous métiers de fonctionnaires civils ; pas seulement les retraites ni l'administration."),
        ("cinq_fois_pensions_tous_statuts_2025", 5 * pensions_total, "HYPOTHÈSE", "Pensions tous statuts y compris militaires ; pas nécessairement postes encore occupés."),
        ("cinq_fois_pensions_civiles_2025", 5 * pensions_civil, "HYPOTHÈSE", "Pensions hors militaires ; tous métiers civils ; aucune prévision de générations."),
        ("sorties_civiles_taux_constant_stock_decroissant", sum(v["sorties_5ans_sans_recrutement"] for v in metrics.values()), "HYPOTHÈSE", "Population de départ égale fin 2024 ; zéro recrutement ; tous métiers civils ; capacité de service non maintenue par construction."),
    ]
    benchmarks = []
    for name, value, status, note in bases:
        for target in [500000, 800000, 1000000, 1500000, 2300000]:
            benchmarks.append({"base": name, "volume_sur_5ans": round(value, 6), "seuil_pp": target,
                               "part_du_volume_necessaire": target / value,
                               "annee": "2027-2031" if status == "HYPOTHÈSE" else "2020-2024",
                               "unite": "personnes_physiques_ou_pensions_selon_base", "statut": status,
                               "source_url": URL_RETRAITES if "pensions" in name else URL_FLUX,
                               "source_localisation": "voir 01b_flux_annuels.csv, 01b_retraites_annuelles.csv et 01b_projection_2027_2031.csv",
                               "note": note + " Une part <=1 n'établit pas une réforme faisable. Seuils PP, non ETP."})
    write_csv("01b_tests_seuils_departs.csv", benchmarks)

    # Contrôles substantiels : agrégats, sources répétées et écarts explicitement conservés.
    def flow2024(branch, stat, indicator):
        return next(r["valeur"] for r in cross if r["groupe"] == branch + "/" + stat and r["indicateur"] == indicator)
    assert sum(flow2024(v, "Ensemble", "sortants") for v in ["FPE", "FPT", "FPH"]) == last_exits
    assert sum(flow2024(v, "Fonctionnaires", "sortants") for v in ["FPE", "FPT", "FPH"]) == last_titulars
    assert pensions_total == thousands(books["retraites"]["Figure 1"]["E17"].value)
    assert pensions_civil == thousands(books["retraites"]["Figure 1"]["E16"].value)
    for col, figurecell in [("D", "E5"), ("G", "E6"), ("J", "E7"), ("M", "E10"), ("P", "E13")]:
        assert thousands(books["comp_retraites"]["F1"][f"{col}6"].value) == thousands(books["retraites"]["Figure 1"][figurecell].value)
    assert sum(next(r["valeur"] for r in movements if r["groupe"] == v and r["indicateur"] == "entrants_autres_versants") for v in ["FPE", "FPT", "FPH"]) == 31942
    assert sum(next(r["valeur"] for r in movements if r["groupe"] == v and r["indicateur"] == "sortants_autres_versants") for v in ["FPE", "FPT", "FPH"]) == 31942

    controls = {
        "date_extraction": DATE, "sorties_civiles_2024_pp": last_exits,
        "sorties_fonctionnaires_2024_pp": last_titulars,
        "sorties_contractuels_2024_pp": flow2024("FPE", "Contractuels", "sortants") + flow2024("FPT", "Contractuels", "sortants") + flow2024("FPH", "Contractuels", "sortants"),
        "part_contractuels_sorties_2024": 261179 / last_exits,
        "cumul_historique_sorties_civiles_2020_2024_occurrences": historical_five,
        "pensions_2025": pensions_total, "pensions_civiles_2025": pensions_civil,
        "pensions_civiles_fonctionnaires_2025": 39146 + 46897 + 22753,
        "pensions_contractuels_2025": 7851 + 9194 + 2142,
        "ecart_stock_flux_pp_par_versant": balances,
        "ecart_contrats_courts_2024_figure5_moins_figure4_pp": thousands(books["flux"]["Fig. 5"]["K6"].value) - thousands(books["flux"]["Fig. 4"]["C10"].value),
        "projections_parametres": metrics,
        "controles": "Agrégats flux, total pensions, cinq croisements pensions vue/compléments, entrées/sorties interversants : conformes.",
        "limite": "Sources DGAFP communes : contrôles croisés internes, pas seconde source statistique indépendante."
    }
    sre_records, sre_reconciliation = read_sre_cache()
    if sre_records:
        write_csv("01b_sre_controle_pensions.csv", sre_records)
        controls["controle_publication_primaire_SRE"] = sre_reconciliation
        controls["limite"] = "Publication primaire distincte SRE vérifiée ; les données administratives SRE/CNRACL sont communes à DGAFP : pas indépendance complète des erreurs."
    (OUT / "01b_controles.json").write_text(json.dumps(controls, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    registry = []
    for key, name in FILES.items():
        registry.append({"id": "DGAFP_01B_" + key.upper(), "titre": name, "institution": "DGAFP-SDessi (sources Insee/Siasp ou caisses de retraite)",
                         "date_publication": "2026-10-15", "url": SOURCE_URLS[key], "date_consultation": DATE,
                         "acces": "integral", "sha256_16": hashlib.sha256((RAW / name).read_bytes()).hexdigest()[:16],
                         "notes": "Empreinte du XLSX local, non de l'archive ZIP. Fichier local intégralement exploité ; fiche primaire et liens d'archive vérifiés sur portail DGAFP. Années observées : flux 2016-2024, retraites 2000-2025, âge 2024 selon fichier."})
    if sre_records:
        registry.append({"id": "SRE_01B_3FP", "titre": "Les pensions de retraite dans les trois fonctions publiques", "institution": "DGFiP/SRE/BFIS et CNRACL",
                         "date_publication": "2026-07", "url": SRE_URL, "date_consultation": DATE, "acces": "integral",
                         "sha256_16": hashlib.sha256(SRE_CACHE.read_bytes()).hexdigest()[:16],
                         "notes": "Page primaire archivée par scripts/fetch.py dans data/cache/01b ; état HTTP 200 et empreinte complète dans manifeste associé. Juillet 2026 indiqué en en-tête ; last-modified HTTP 21/10/2026. Écart FPE avec DGAFP entièrement expliqué par La Poste et Orange."})
    (ROOT / "sources").mkdir(exist_ok=True)
    with (ROOT / "sources/registre_01b.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(registry[0]))
        writer.writeheader()
        writer.writerows(registry)
    for book in books.values():
        book.close()
    print(json.dumps(controls, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
