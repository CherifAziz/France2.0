#!/usr/bin/env python3
"""Inspection bornée d'un XLSX, y compris dimensions gonflées par le formatage."""
import argparse
import json
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("file")
parser.add_argument("--sheet")
parser.add_argument("--range", default="A1:AZ85")
args = parser.parse_args()
book = load_workbook(args.file, read_only=True, data_only=True)
if not args.sheet:
    print(json.dumps({s.title: [s.max_row, s.max_column] for s in book}, ensure_ascii=False))
else:
    sheet = book[args.sheet]
    c1, r1, c2, r2 = range_boundaries(args.range)
    for row_number, row in enumerate(sheet.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2), r1):
        vals = {f"{get_column_letter(c1 + i)}{row_number}": cell.value
                for i, cell in enumerate(row) if cell.value is not None}
        if vals:
            print(json.dumps(vals, ensure_ascii=False, default=str))
book.close()
