#!/usr/bin/env python3
"""Livrables complémentaires de la phase 6 et chapitres générés du rapport.

Écrit : data/reformes/14_matrice_risques.csv, livrables/bibliographie.csv,
livrables/tableaux_livrables.xlsx (réformes, risques, base comparative, seuils, sensibilités),
figures/14_emploi_s13_international.svg, figures/14_scenarios_seuils.svg,
rapport/11_catalogue.md, rapport/16_bibliographie.md, rapport/17_annexes.md.
"""
import csv
import glob
import importlib.util
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SC = ["prudent", "intermediaire", "ambitieux", "radical"]
LAB = {"prudent": "Prudent", "intermediaire": "Intermédiaire", "ambitieux": "Ambitieux", "radical": "Radical"}
CHANTIERS = {"01": "Emploi public en France", "01b": "Départs et retraites", "01c": "Coûts et pensions", "02": "Métiers",
             "03": "Base internationale", "04_controle": "Contrôle des données héritées", "04a": "Fiscalité et emploi",
             "04b": "Éducation, police, justice", "04c": "Santé et protection sociale", "04d": "Territoires et fonctions support",
             "05": "Réformes étrangères", "06": "Structures et précédents français", "07": "Simplification",
             "08": "IA en France", "09": "Preuves internationales sur l'IA", "10": "Audit des propositions"}


def read(path, enc="utf-8"):
    with open(ROOT / path, encoding=enc) as f:
        return list(csv.DictReader(f))


def fr(x, d=0):
    s = f"{x:,.{d}f}".replace(",", " ").replace(".", ",")
    return s.replace("-", "−")


def load_fiches():
    out = []
    for name in ("11_fiches_etat.py", "11_fiches_local_social.py"):
        spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / name)
        m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(m)
        out += m.FICHES
    return {f["id"]: f for f in out}


def first_sentence(text):
    text = re.sub(r"\*\*\[[^\]]+\]\*\*\s*", "", text).replace("**", "")
    return re.split(r"(?<=[.;])\s", text, 1)[0].rstrip(" ;")


TRANSVERSES = [
    ("T1", "Les gisements de postes ne se matérialisent pas", "élevée", "élevé", "24 fiches sur 30 à confiance faible ; 66 % des postes du scénario intermédiaire (note 13, O1)", "Pilotes mesurés avant toute généralisation ; budget voté sur économies constatées, pas annoncées", "Postes et crédits à périmètre constant, 24-36 mois"),
    ("T2", "Rebond des effectifs après la contraction", "moyenne", "élevé", "Royaume-Uni +36,2 % après 2016 ; Irlande revenue à 317 495 WTE en 2017 (note 05)", "Suppression des missions et pas seulement des postes ; plafonds d'emplois pluriannuels", "Effectifs et recours aux contractuels 5 ans après"),
    ("T3", "Substitution par des prestataires ou des contractuels", "moyenne", "moyen", "État territorial 2012-2020 compensé par des contrats courts ; dépenses externes néo-zélandaises +7,18 % (notes 05-06)", "Suivi consolidé masse salariale + prestations + contrats courts", "Dépenses de prestations et ETP de contractuels"),
    ("T4", "Dépassement des investissements numériques", "moyenne", "élevé", "Phénix : 937,5 M CAD en 2024-25, successeur > 4,2 Md CAD ; Sirhen ≈ 480 M€ (notes 05, 08)", "Projets découpés, réversibles, double run, avis DINUM ; budget de risque explicite", "Écart coût réel / devis ; incidents de paie"),
    ("T5", "Dégradation du service et de l'accès aux droits", "moyenne", "élevé", "7 % des 16-74 ans en illectronisme ; Robodebt, Toeslagen (notes 08-09)", "Accueil humain garanti ; indicateurs de délai et de recours publiés ; mode dégradé", "Délais, réclamations, recours fondés, non-recours"),
    ("T6", "Risque juridique des décisions automatisées", "moyenne", "moyen", "RGPD art. 22, CRPA, AI Act (annexe III au 2/12/2027) ; texte opposable non relu (note 08)", "Analyse d'impact par usage ; contrôle humain doté du pouvoir de corriger", "Contentieux perdus ; décisions corrigées"),
    ("T7", "Pertes de recettes par moindre contrôle", "faible", "moyen", "Fonctions fiscales d'audit et de recouvrement maintenues ; R03 en redéploiement (note 04a)", "Ne pas réduire les fonctions de contrôle à rendement net positif", "Rendement net du contrôle fiscal"),
    ("T8", "Départs moins nombreux que prévu", "faible", "moyen", "Taux de départ homogène par versant ; pyramide des âges des filières non publiée (note 13, O8)", "Calendrier ajusté aux départs réels ; pas de licenciement", "Départs constatés par vivier"),
    ("T9", "Double compte avec la trajectoire de référence", "moyenne", "faible", "Tendance DGFiP −15 % en 10 ans ; PLF 2027, foncière (notes 11, 13)", "Référence publiée avant chaque réforme ; registre des viviers tenu à jour", "Économies mesurées contre la référence publiée"),
]


def matrice(fiches, res_inter):
    rows = []
    for t in TRANSVERSES:
        rows.append(dict(zip(["id", "risque", "probabilite", "impact", "preuve", "attenuation", "indicateur"], t)) | {"type": "transversal", "reforme": "", "postes_2036_intermediaire": ""})
    proba = {"faible": "élevée", "moyen": "moyenne", "élevé": "faible"}
    for rid, f in fiches.items():
        p = res_inter.get(rid, {})
        rows.append({"id": rid, "type": "réforme", "reforme": f["titre"], "risque": first_sentence(f["I"]),
                     "probabilite": proba[f["confiance"]], "impact": "moyen", "preuve": "Confiance " + f["confiance"] + " : " + first_sentence(f["L"]),
                     "attenuation": first_sentence(f["J"]), "indicateur": first_sentence(f["V"]),
                     "postes_2036_intermediaire": round(float(p.get("etp_postes_en_moins", 0)))})
    fields = ["id", "type", "reforme", "risque", "probabilite", "impact", "preuve", "attenuation", "indicateur", "postes_2036_intermediaire"]
    with open(ROOT / "data/reformes/14_matrice_risques.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    return rows, fields


def bibliographie():
    rows = []
    for path in sorted(glob.glob(str(ROOT / "sources/registre_*.csv"))):
        key = Path(path).stem.replace("registre_", "")
        for r in read(Path(path).relative_to(ROOT), "utf-8-sig"):
            rows.append({"chantier": key, "id": r.get("id", ""), "titre": r.get("titre", ""), "institution": r.get("institution", ""),
                         "date_publication": r.get("date_publication", ""), "url": r.get("url", ""), "date_consultation": r.get("date_consultation", ""),
                         "acces": r.get("acces", "")})
    seen, uniq = {}, []
    for r in rows:
        k = r["url"] or r["id"]
        if k in seen:
            seen[k]["chantier"] += ", " + r["chantier"]
            continue
        seen[k] = r
        uniq.append(r)
    with open(ROOT / "livrables/bibliographie.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(uniq[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(uniq)
    acc = {}
    for r in uniq:
        acc[r["acces"] or "non précisé"] = acc.get(r["acces"] or "non précisé", 0) + 1
    md = ["# 16. Bibliographie", "",
          f"Bibliographie fusionnée des {len(glob.glob(str(ROOT / 'sources/registre_*.csv')))} registres de sources de la mission : **{len(rows)} entrées, {len(uniq)} références distinctes** après dédoublonnage par URL. "
          "Le statut d'accès indique si le document a été effectivement lu (« integral »), lu en partie, ou s'il est resté inaccessible : un document inaccessible n'est jamais présenté comme lu. "
          "Répartition : " + " ; ".join(f"{k} : {v}" for k, v in sorted(acc.items(), key=lambda x: -x[1])) + ". "
          "Fichier complet avec dates de consultation : `livrables/bibliographie.csv` ; empreintes et archives : `sources/registre_*.csv`.", ""]
    by = {}
    for r in uniq:
        by.setdefault(r["chantier"].split(",")[0], []).append(r)
    for key in sorted(by, key=lambda k: list(CHANTIERS).index(k) if k in CHANTIERS else 99):
        md += [f"## {key} — {CHANTIERS.get(key, key)}", ""]
        for r in sorted(by[key], key=lambda r: (r["institution"], r["titre"])):
            date = f", {r['date_publication']}" if r["date_publication"] else ""
            title = r["titre"].replace("|", "/")
            link = f"[{title}]({r['url']})" if r["url"].startswith("http") else title
            md.append(f"- {r['institution']}{date}. {link}. Consulté le {r['date_consultation']} ; accès : {r['acces'] or 'non précisé'}.")
        md.append("")
    (ROOT / "rapport/16_bibliographie.md").write_text("\n".join(md), encoding="utf-8", newline="\n")
    return uniq


def svg_bars(path, title, subtitle, items, unit, highlight="FRA", ref=None, decimals=1):
    W, left, top, row = 820, 200, 64, 22
    H = top + row * len(items) + 40
    vmax = max([v for _, v in items] + [v for v, _ in (ref or [])]) * 1.15
    x = lambda v: left + v / vmax * (W - left - 70)
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" font-family="Segoe UI, Arial, sans-serif" font-size="12">',
           f'<rect width="{W}" height="{H}" fill="#ffffff"/>',
           f'<text x="16" y="26" font-size="16" font-weight="600" fill="#17232b">{title}</text>',
           f'<text x="16" y="46" fill="#53616b">{subtitle}</text>']
    for i, (lab, v) in enumerate(items):
        y = top + i * row
        color = "#c0392b" if lab.startswith(highlight) else ("#7f8c8d" if lab.startswith("Moyenne") else "#245574")
        out.append(f'<text x="{left - 8}" y="{y + 14}" text-anchor="end" fill="#17232b">{lab}</text>')
        out.append(f'<rect x="{left}" y="{y + 3}" width="{x(v) - left:.1f}" height="15" fill="{color}"/>')
        clash = any(x(v) < x(r) < x(v) + 70 for r, _ in (ref or []))
        if clash:  # étiquette dans la barre si elle croiserait une ligne de seuil
            out.append(f'<text x="{x(v) - 5:.1f}" y="{y + 15}" text-anchor="end" fill="#ffffff">{fr(v, decimals)}{unit}</text>')
        else:
            out.append(f'<text x="{x(v) + 5:.1f}" y="{y + 15}" fill="#53616b">{fr(v, decimals)}{unit}</text>')
    if ref:
        for v, lab in ref:
            out.append(f'<line x1="{x(v):.1f}" y1="{top}" x2="{x(v):.1f}" y2="{top + row * len(items) + 4}" stroke="#c0392b" stroke-dasharray="4 3"/>')
            out.append(f'<text x="{x(v):.1f}" y="{top + row * len(items) + 16}" text-anchor="middle" fill="#c0392b" font-size="11">{lab}</text>')
    out.append("</svg>")
    path.write_text("\n".join(out) + "\n", encoding="utf-8", newline="\n")


NOMS = {"SWE": "Suède", "DNK": "Danemark", "FIN": "Finlande", "EST": "Estonie", "FRA": "France", "CAN": "Canada", "IRL": "Irlande",
        "BEL": "Belgique", "GBR": "Royaume-Uni", "POL": "Pologne", "AUT": "Autriche", "ESP": "Espagne", "AUS": "Australie", "USA": "États-Unis",
        "ITA": "Italie", "NLD": "Pays-Bas", "DEU": "Allemagne", "OECD": "Moyenne OCDE", "JPN": "Japon", "KOR": "Corée", "NZL": "Nouvelle-Zélande", "CZE": "Tchéquie"}


def figures(synth, seuils):
    base = read("data/international/03_base_comparative.csv", "utf-8-sig")
    vals = {r["pays_iso3"]: float(r["valeur"]) for r in base if r["indicateur"] == "emploi_s13_pct_emploi" and r["annee"] == "2023"
            and r["source_id"] == "03_oecd_gag2025" and r["pays_iso3"] in NOMS and r["valeur"]}
    items = sorted(((("Moyenne OCDE" if k == "OECD" else NOMS[k]), v) for k, v in vals.items()), key=lambda x: -x[1])
    items = [(("FRA " if n == "France" else "") + n if False else n, v) for n, v in items]
    svg_bars(ROOT / "figures/14_emploi_s13_international.svg", "Emploi des administrations publiques en % de l'emploi total, 2023",
             "OCDE, Panorama des administrations publiques 2025 (StatLink) ; champ S13, pas une mesure de l'emploi administratif", items, " %", highlight="France")
    s = seuils["500000"]
    items2 = [(f"Catalogue {LAB[sc].lower()}", float(synth[sc]["postes_en_moins_10ans"])) for sc in SC]
    items2 += [("Radical, hyp. favorables", float(s["radical_favorable_2036"])),
               ("Capacité retraites 10 ans", float(s["capacite_retraites_10ans"])),
               ("Capacité toutes sorties 10 ans", float(s["capacite_toutes_sorties_10ans"]))]
    svg_bars(ROOT / "figures/14_scenarios_seuils.svg", "Postes en moins en 2036 et capacités de départs, comparés aux seuils",
             "Moteur financier et capacités B × (1 − (1 − r)^10) des sept viviers administratifs (notes 11 à 13)", items2, "", highlight="§",
             ref=[(500000, "500 000"), (800000, "800 000")], decimals=0)


def classeur(rows_risk, risk_fields, synth):
    wb = Workbook()
    head = lambda ws: [setattr(c, "font", Font(bold=True)) or setattr(c, "fill", PatternFill("solid", fgColor="DCE6EF")) for c in ws[1]]
    ws = wb.active
    ws.title = "Reformes"
    cat = read("data/reformes/catalogue_reformes.csv")
    res = {sc: {r["id"]: r for r in read(f"modele/resultats/reformes_{sc}.csv")} for sc in SC}
    cols = ["id", "intitule", "vivier_id", "confiance", "annee_debut", "duree_ans"]
    ws.append(cols + [f"postes_2036_{s}" for s in SC] + [f"net_2036_meur_{s}" for s in SC] + [f"cumul_net_10ans_meur_{s}" for s in SC] + [f"invest_meur_{s}" for s in SC] + ["fiche"])
    for r in cat:
        ws.append([r[c] if c not in ("annee_debut", "duree_ans") else int(float(r[c])) for c in cols]
                  + [round(float(res[s][r["id"]]["etp_postes_en_moins"])) for s in SC]
                  + [round(float(res[s][r["id"]]["net_annee_finale_meur"]), 1) for s in SC]
                  + [round(float(res[s][r["id"]]["cumul_net_10ans"]), 1) for s in SC]
                  + [float(r[f"invest_meur_{s}"]) for s in SC] + [r["fiche"]])
    ws.append(["Total", "", "", "", "", ""] + [round(float(synth[s]["postes_en_moins_10ans"])) for s in SC] + [round(float(synth[s]["eco_nette_2036_meur"]), 1) for s in SC]
              + [round(float(synth[s]["cumul_net_10ans_meur"]), 1) for s in SC] + [round(float(synth[s]["invest_horizon_meur"]), 1) for s in SC])
    head(ws)
    for name, path in [("Matrice des risques", "data/reformes/14_matrice_risques.csv"), ("Seuils", "data/reformes/12_seuils.csv"),
                       ("Sensibilites", "data/reformes/12_sensibilites.csv"), ("Equipe B", "data/reformes/13_equipe_b.csv"),
                       ("Base comparative", "data/international/03_base_comparative.csv"), ("Dictionnaire base", "data/international/03_dictionnaire_variables.csv")]:
        data = read(path, "utf-8-sig")
        sh = wb.create_sheet(name)
        sh.append(list(data[0]))
        for r in data:
            sh.append([_num(v) for v in r.values()])
        head(sh)
    for sh in wb.worksheets:
        sh.freeze_panes = "A2"
    wb.save(ROOT / "livrables/tableaux_livrables.xlsx")


def _num(v):
    try:
        return float(v) if re.fullmatch(r"-?\d+(\.\d+)?(e-?\d+)?", v or "") else v
    except ValueError:
        return v


def chapitre_11(fiches, synth):
    res = {sc: {r["id"]: r for r in read(f"modele/resultats/reformes_{sc}.csv")} for sc in SC}
    md = ["# 11. Catalogue détaillé des réformes", "",
          "Trente fiches suivent la structure A à M du cahier des charges, avec un indicateur de validation à deux ou trois ans. "
          "Chaque fait cité provient d'une note contre-vérifiée ; chaque taux de réduction est une hypothèse rattachée à un précédent observé, "
          "et les économies de la rubrique G sont celles du moteur financier (chapitre 13). La méthode des viviers, les hypothèses nationales "
          "et les conventions sont détaillées dans la [note 11](../recherche/11_catalogue_reformes.md).", "",
          "## 11.1 Vue d'ensemble", "",
          "| Fiche | Confiance | Postes en moins 2036 (P / I / A / R) | Économie nette 2036, M€ (P / I / A / R) |", "|---|---|---|---|"]
    for rid, f in fiches.items():
        p = " / ".join(fr(float(res[s][rid]["etp_postes_en_moins"])) for s in SC)
        n = " / ".join(fr(float(res[s][rid]["net_annee_finale_meur"])) for s in SC)
        md.append(f"| {rid} — {f['titre']} | {f['confiance']} | {p} | {n} |")
    md.append(f"| **Total** | | **{' / '.join(fr(float(synth[s]['postes_en_moins_10ans'])) for s in SC)}** | **{' / '.join(fr(float(synth[s]['eco_nette_2036_meur'])) for s in SC)}** |")
    md += ["", "P = prudent, I = intermédiaire, A = ambitieux, R = radical. Totaux non pondérés ; la contre-expertise (chapitre 14) les pondère par la confiance.", "",
           "## 11.2 Les trente fiches", ""]
    cat = {r["id"]: r for r in read("data/reformes/catalogue_reformes.csv")}
    for rid in fiches:
        text = (ROOT / cat[rid]["fiche"]).read_text(encoding="utf-8")
        text = re.sub(r"^# Réforme ", "### Réforme ", text, flags=re.M)
        text = text.replace("](../", "](../recherche/").replace("](../recherche/../", "](../")
        md += [text.strip(), ""]
    (ROOT / "rapport/11_catalogue.md").write_text("\n".join(md), encoding="utf-8", newline="\n")


def chapitre_17():
    md = ["# 17. Annexes statistiques et calculs reproductibles", "",
          "Toutes les données, scripts et contrôles de la mission sont versionnés dans le dépôt. Les documents lourds (PDF de rapports) sont conservés hors Git dans `data/cache/` ; leurs URL et empreintes figurent dans les registres, et les scripts `restore_lot*_sources.py` les restaurent.", "",
          "## 17.1 Livrables complémentaires", "",
          "| Livrable | Fichier |", "|---|---|",
          "| Base comparative internationale structurée (3 859 lignes, 44 pays) | `data/international/03_base_comparative.csv`, dictionnaire `03_dictionnaire_variables.csv` ; feuille « Base comparative » de `livrables/tableaux_livrables.xlsx` |",
          "| Tableau exhaustif des réformes et de leurs économies | `data/reformes/catalogue_reformes.csv`, `modele/resultats/reformes_*.csv` ; feuille « Reformes » de `livrables/tableaux_livrables.xlsx` |",
          "| Modèle financier modifiable | `modele/modele_financier.xlsx` (formules, hypothèses en cellules jaunes) et moteur `modele/modele.py` |",
          "| Matrice des risques et incertitudes | `data/reformes/14_matrice_risques.csv` ; feuille « Matrice des risques » |",
          "| Synthèse de 10 pages | `livrables/synthese.html`, `.docx`, `.pdf` |",
          "| Bibliographie fusionnée | `livrables/bibliographie.csv` |", "",
          "## 17.2 Notes de recherche détaillées", "", "| Note | Objet |", "|---|---|"]
    for p in sorted((ROOT / "recherche").glob("[0-1]*.md")):
        first = p.read_text(encoding="utf-8").splitlines()[0].lstrip("# ").strip()
        md.append(f"| [{p.name}](../recherche/{p.name}) | {first} |")
    md += ["", "## 17.3 Contrôles et contre-vérifications", "",
           "| Contrôle | Chiffres | Résultat | Fichier |", "|---|---:|---|---|",
           "| Lot 1 (emploi, départs, coûts, métiers, base internationale) | 100 | 98 confirmés, 2 reclassés en estimation | `data/risques/verification_40_chiffres_lot1.csv` |",
           "| Lot 2 (comparaisons par fonction) | 63 | 61 conformes, 2 après correction | `data/risques/verification_chiffres_lot2.csv` |",
           "| Lot 3 (réformes étrangères, structures, simplification, IA, propositions) | 98 | aucun chiffre contredit ; 3 corrections de forme | `data/risques/verification_chiffres_lot3.csv` |",
           "| Phase 5 (catalogue, hypothèses, scénarios) | 47 | aucun écart ; 2 réserves | `data/risques/verification_phase5.csv` |", "",
           "## 17.4 Reproduction complète", "", "```powershell",
           "python -I scripts/11_catalogue.py          # catalogue, viviers, moteur, fiches",
           "python -I scripts/12_sensibilites.py       # sensibilités et seuils",
           "python -I scripts/12_classeur.py           # classeur Excel",
           "python -I scripts/00_verification_phase5.py",
           "python -I scripts/13_equipe_b.py",
           "python -I scripts/14_livrables.py          # matrice, bibliographie, figures, chapitres générés",
           "python scripts/build_rapport.py --pdf      # rapport et synthèse (Markdown, python-docx, Playwright)",
           "python -I -m unittest discover -s tests -p \"test_modele.py\"", "```", "",
           "## 17.5 Dictionnaire des principales variables du modèle", "",
           "| Variable | Définition | Unité |", "|---|---|---|",
           "| `etp_base` | Effectif initial du vivier | ETP |",
           "| `taux_departs_annuel` | Départs mobilisables par an, rapportés au stock restant | fraction |",
           "| `plafond_reduction_{scénario}` | Part maximale du vivier réductible | fraction |",
           "| `taux_{scénario}` | Part du reliquat du vivier visée par une réforme | fraction |",
           "| `part_non_remplacement`, `part_redeploiement`… | Répartition des mouvements | fraction, somme = 1 |",
           "| `invest_meur_{scénario}` | Investissement total de la réforme | M€ 2026 |",
           "| `cout_recurrent_meur_{scénario}` | Coût annuel supplémentaire à terme | M€/an |",
           "| `eco_fonct_meur_{scénario}`, `effet_recettes_meur_{scénario}` | Économie de fonctionnement, recettes | M€/an |",
           "| `cout_entree_hors_pension`, `cout_moyen_hors_pension` | Coût employeur annuel hors pension | k€/ETP |",
           "| `gvt_individuel_reel`, `taux_conservation_consolide` | Progression du coût d'une cohorte ; part conservée en lecture APU | fraction |", ""]
    (ROOT / "rapport/17_annexes.md").write_text("\n".join(md), encoding="utf-8", newline="\n")


def main():
    (ROOT / "rapport").mkdir(exist_ok=True)
    (ROOT / "livrables").mkdir(exist_ok=True)
    fiches = load_fiches()
    synth = {r["scenario"]: r for r in read("modele/resultats/synthese_scenarios.csv")}
    seuils = {r["seuil"]: r for r in read("data/reformes/12_seuils.csv")}
    res_inter = {r["id"]: r for r in read("modele/resultats/reformes_intermediaire.csv")}
    rows, fields = matrice(fiches, res_inter)
    biblio = bibliographie()
    figures(synth, seuils)
    classeur(rows, fields, synth)
    chapitre_11(fiches, synth)
    chapitre_17()
    print(f"{len(rows)} risques ; {len(biblio)} références ; chapitres 11, 16, 17 ; 2 figures ; tableaux_livrables.xlsx")


if __name__ == "__main__":
    main()
