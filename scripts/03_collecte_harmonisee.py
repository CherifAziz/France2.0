#!/usr/bin/env python3
"""Collecte descriptive S13 et comptes publics, millésimes 2023 et 2024.

Python 3 + requests. Ne modifie pas les données de la session précédente.
Sans --refresh, relit les instantanés archivés. Aucune imputation des absences.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import itertools
import json
import math
import time
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urlencode

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/03"
OUT = ROOT / "data/international"
YEARS = (2023, 2024)
DATE = "2026-10-28"
COUNTRIES = {
    "AUS": ("Australie", "AU"), "AUT": ("Autriche", "AT"), "BEL": ("Belgique", "BE"),
    "BGR": ("Bulgarie", "BG"), "CAN": ("Canada", "CA"), "CHE": ("Suisse", "CH"),
    "CHL": ("Chili", "CL"), "COL": ("Colombie", "CO"), "CRI": ("Costa Rica", "CR"),
    "CYP": ("Chypre", "CY"), "CZE": ("Tchéquie", "CZ"), "DEU": ("Allemagne", "DE"),
    "DNK": ("Danemark", "DK"), "ESP": ("Espagne", "ES"), "EST": ("Estonie", "EE"),
    "FIN": ("Finlande", "FI"), "FRA": ("France", "FR"), "GBR": ("Royaume-Uni", "UK"),
    "GRC": ("Grèce", "EL"), "HRV": ("Croatie", "HR"), "HUN": ("Hongrie", "HU"),
    "IRL": ("Irlande", "IE"), "ISL": ("Islande", "IS"), "ISR": ("Israël", "IL"),
    "ITA": ("Italie", "IT"), "JPN": ("Japon", "JP"), "KOR": ("Corée", "KR"),
    "LTU": ("Lituanie", "LT"), "LUX": ("Luxembourg", "LU"), "LVA": ("Lettonie", "LV"),
    "MEX": ("Mexique", "MX"), "MLT": ("Malte", "MT"), "NLD": ("Pays-Bas", "NL"),
    "NOR": ("Norvège", "NO"), "NZL": ("Nouvelle-Zélande", "NZ"), "POL": ("Pologne", "PL"),
    "PRT": ("Portugal", "PT"), "ROU": ("Roumanie", "RO"), "SGP": ("Singapour", "SG"),
    "SVK": ("Slovaquie", "SK"), "SVN": ("Slovénie", "SI"), "SWE": ("Suède", "SE"),
    "TUR": ("Turquie", "TR"), "USA": ("États-Unis", "US"),
}
EURO_GEOGRAPHIES = "AT BE BG CH CY CZ DE DK EE EL ES EU27_2020 FI FR HR HU IE IS IT LT LU LV MT NL NO PL PT RO SE SI SK".split()
ISO2_TO_ISO3 = {value[1]: key for key, value in COUNTRIES.items()}
SCHEMAS = {
    "emploi_s13_personnes": ("personnes", "Personnes employées par les unités institutionnelles S13, concept intérieur ; aucun classement par métier"),
    "emploi_total_interieur_personnes": ("personnes", "Emploi total de l'économie S1, concept intérieur, tous statuts"),
    "population_personnes": ("personnes", "Population totale annuelle de la comptabilité nationale, champ national de la source"),
    "emploi_s13_pct_emploi": ("% emploi intérieur", "100 × emploi S13 / emploi total intérieur, personnes, même pays et année"),
    "emploi_s13_pour_1000_habitants": ("personnes pour 1 000 habitants", "1 000 × emploi S13 intérieur / population annuelle ; sensibilité aux travailleurs frontaliers"),
    "D1_pct_pib": ("% PIB", "Rémunération des salariés S13, salaires et cotisations employeur, toutes fonctions"),
    "TE_pct_pib": ("% PIB", "Dépenses totales consolidées S13, toutes fonctions"),
    "P2_pct_pib": ("% PIB", "Consommation intermédiaire S13 ; ne mesure pas la seule administration"),
    "D632_pct_pib": ("% PIB", "Transferts sociaux en nature, production marchande achetée par S13"),
    "P2_D632_pct_pib": ("% PIB", "P2 + D632 ; achats de biens/services et prestations en nature, pas une mesure stricte d'externalisation administrative"),
    "D1_P2_D632_pct_pib": ("% PIB", "Sous-total D1+P2+D632 ; ne comprend pas tous les coûts de production, notamment consommation de capital fixe"),
    "GF01_hors_interets_pct_pib": ("% PIB", "Non calculé : D41 ventilé par fonction non trouvé ; GF0107 ne doit pas être assimilé intégralement aux intérêts"),
}


def code(value):
    return (value or "").split(":", 1)[0].strip()


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"Table vide : {path}")
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def archive(name, url, refresh=False, accept=None):
    path = RAW / name
    meta_path = path.with_suffix(path.suffix + ".meta.json")
    if refresh or not path.exists():
        headers = {"User-Agent": "France20-independent-research/1.0"}
        if accept:
            headers["Accept"] = accept
        for attempt in range(3):
            response = requests.get(url, headers=headers, timeout=45)
            if response.status_code not in (429, 502, 503, 504) or attempt == 2:
                break
            time.sleep(2 * (attempt + 1))
        response.raise_for_status()
        body = response.content
        if not body:
            raise ValueError(f"Réponse vide : {url}")
        if name.endswith(".csv") and b"OBS_VALUE" not in body[:6000]:
            raise ValueError(f"Réponse sans colonne SDMX OBS_VALUE : {url}")
        if name.endswith(".json"):
            data = response.json()
            if data.get("class") != "dataset":
                raise ValueError(f"JSON sans dataset : {url}")
        if name.endswith(".xml"):
            ET.fromstring(body)
        RAW.mkdir(parents=True, exist_ok=True)
        path.write_bytes(body)
        meta_path.write_text(json.dumps({"url": url, "url_finale": response.url, "date_consultation": DATE,
                             "http_status": response.status_code, "content_type": response.headers.get("Content-Type", ""),
                             "server_date": response.headers.get("Date", ""), "etag": response.headers.get("ETag", ""),
                             "sha256": hashlib.sha256(body).hexdigest()}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    body = path.read_bytes()
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if hashlib.sha256(body).hexdigest() != meta["sha256"]:
        raise ValueError(f"Empreinte différente du manifeste : {path}")
    return path, meta


def structure(flow, version, refresh):
    url = f"https://sdmx.oecd.org/public/rest/v1/dataflow/OECD.SDD.NAD/{flow}/{version}?references=datastructure"
    path, meta = archive("structure_" + flow.split("@")[-1] + ".xml", url, refresh)
    root = ET.fromstring(path.read_bytes())
    dimensions = sorted([(int(x.attrib["position"]), x.attrib["id"]) for x in root.iter()
                         if x.tag.split("}")[-1] == "Dimension" and "position" in x.attrib])
    return [name for _, name in dimensions], path, meta


def oecd_source(source_id, flow, version, filters, refresh):
    dimensions, structure_path, structure_meta = structure(flow, version, refresh)
    unknown = set(filters) - set(dimensions)
    if unknown:
        raise ValueError(f"Dimensions inconnues pour {flow}: {unknown}")
    key = ".".join(filters.get(d, "") for d in dimensions)
    url = f"https://sdmx.oecd.org/public/rest/v1/data/OECD.SDD.NAD,{flow},{version}/{key}?startPeriod=2023&endPeriod=2024"
    path, meta = archive(source_id + ".csv", url, refresh, "application/vnd.sdmx.data+csv; charset=utf-8; labels=both")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    rows = [{k.split(":", 1)[0]: v for k, v in row.items()} for row in rows]
    return {"id": source_id, "institution": "OCDE", "rows": rows, "path": path, "meta": meta,
            "structure_path": structure_path, "structure_meta": structure_meta, "dimensions": dimensions}


def euro_source(source_id, dataset, filters, refresh):
    params = [("lang", "en"), ("sinceTimePeriod", "2023"), ("untilTimePeriod", "2024")]
    for key, values in filters.items():
        params.extend((key, value) for value in values)
    url = "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/" + dataset + "?" + urlencode(params)
    path, meta = archive(source_id + ".json", url, refresh)
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    ids = data["id"]
    categories = []
    for dimension in ids:
        index = data["dimension"][dimension]["category"]["index"]
        categories.append([x[0] for x in sorted(index.items(), key=lambda x: x[1])] if isinstance(index, dict) else index)
    values, statuses = data.get("value", {}), data.get("status", {})
    rows = []
    for index, combo in enumerate(itertools.product(*categories)):
        value = values.get(str(index)) if isinstance(values, dict) else values[index]
        if value is None:
            continue
        status = statuses.get(str(index), "") if isinstance(statuses, dict) else statuses[index] if index < len(statuses) else ""
        rows.append({**dict(zip(ids, combo)), "value": value, "flag": status})
    return {"id": source_id, "institution": "Eurostat", "rows": rows, "path": path, "meta": meta,
            "updated": data.get("updated", ""), "dimensions": ids}


def oecd_value(row):
    value = row.get("OBS_VALUE", "")
    if not value:
        return None
    exponent = code(row.get("UNIT_MULT", ""))
    if exponent == "":
        raise ValueError("UNIT_MULT manquant : pas de conversion implicite")
    result = float(value) * 10 ** int(exponent)
    return result if math.isfinite(result) else None


def unique(rows, **filters):
    selected = [r for r in rows if all(code(r.get(k, "")) == str(v) for k, v in filters.items())]
    if len(selected) > 1:
        raise ValueError(f"Clé insuffisamment sélective : {filters}, {len(selected)} lignes")
    return selected[0] if selected else None


def row_output(country, year, indicator, source, value, flag="", note="", formula="", inputs=()):
    unit, scope = SCHEMAS[indicator]
    missing = value is None
    sources = [source, *inputs]
    urls = list(dict.fromkeys(s["meta"]["url"] for s in sources))
    paths = list(dict.fromkeys(str(s["path"].relative_to(ROOT)).replace("\\", "/") for s in sources))
    special_names = {"EU27_2020": "Union européenne à 27", "OECD": "Moyenne OCDE publiée (composition propre à la source)"}
    return dict(pays_iso3=country, pays=COUNTRIES.get(country, (special_names.get(country, country), ""))[0],
                indicateur=indicator, valeur="" if missing else round(value, 9), unite=unit, annee=year,
                champ=scope, secteur="S13" if "s13" in indicator or "pct_pib" in indicator else "S1",
                origine=source["institution"], source_id=source["id"], source_url=" ; ".join(urls),
                edition_source=source.get("edition", "API consultée le " + DATE),
                source_localisation=" ; ".join(paths)+f" ; pays={country} ; année={year}",
                flag=flag or ("ABSENT_EXTRACTION" if missing else ""),
                statut="[NON TROUVÉ]" if missing else "[ESTIMATION]" if formula else "[FAIT]",
                formule=formula, note=note,
                date_consultation=DATE)


def panorama_source(refresh):
    path, meta = archive("03_oecd_gag2025_fig13_1.xlsx", "https://stat.link/8htodn", refresh)
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["g13-1"]
    if sheet["C31"].value != 2023 or sheet["D31"].value != 2019:
        raise ValueError("Structure inattendue du StatLink : vérifier manuellement les années")
    rows = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=32, max_col=4, values_only=True), 32):
        _, country, value, _ = cells
        if not country or value is None:
            continue
        if country not in COUNTRIES and country != "OECD":
            raise ValueError(f"Pays StatLink non prévu : {country}")
        year = 2022 if country in {"ISR", "NOR", "PRT", "CHE"} else 2021 if country == "CRI" else 2023
        rows.append({"country": country, "year": year, "value": float(value), "cell": f"g13-1!C{row_number}"})
    book.close()
    return {"id": "03_oecd_gag2025", "institution": "OCDE", "rows": rows, "path": path, "meta": meta,
            "publication": "2025-06-19", "edition": "Government at a Glance 2025 ; StatLink version 1 du 19 juin 2025"}


def financial_reconciliation(output):
    index = {(r["pays_iso3"], r["annee"], r["indicateur"]): r for r in output
             if r["source_id"] == "03_oecd_fiscal" and r["valeur"] != ""}
    rows = []
    for item in output:
        if item["source_id"] != "03_eurostat_fiscal" or item["valeur"] == "":
            continue
        other = index.get((item["pays_iso3"], item["annee"], item["indicateur"]))
        if not other:
            continue
        rows.append(dict(pays_iso3=item["pays_iso3"], annee=item["annee"], indicateur=item["indicateur"],
                         valeur_eurostat=item["valeur"], valeur_ocde_calculee=other["valeur"],
                         ecart_points=float(other["valeur"])-float(item["valeur"]),
                         source_url=item["source_url"]+" ; "+other["source_url"],
                         source_localisation=item["source_localisation"]+" ; "+other["source_localisation"],
                         statut="[ESTIMATION]", note="Données nationales souvent communes ; ce rapprochement n'est pas une source indépendante"))
    write_csv(OUT / "03_rapprochement_finances.csv", rows)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--refresh", action="store_true", help="Rafraîchir les instantanés officiels")
    args = parser.parse_args()
    countries = "+".join(COUNTRIES)
    common = {"FREQ": "A", "REF_AREA": countries, "TRANSFORMATION": "N"}
    jobs = [
        lambda: oecd_source("03_oecd_emp_gov", "DSD_NASEC10@DF_TABLE12_EMPDC", "1.1", {**common, "SECTOR": "S13", "TRANSACTION": "EMP", "UNIT_MEASURE": "PS"}, args.refresh),
        lambda: oecd_source("03_oecd_emp_sector", "DSD_NAMAIN10@DF_TABLE3_EMP_SECTOR", "2.0", {**common, "SECTOR": "S1+S13", "TRANSACTION": "EMP", "UNIT_MEASURE": "PS"}, args.refresh),
        lambda: oecd_source("03_oecd_emp_total", "DSD_NAMAIN10@DF_TABLE3_EMPDC", "2.0", {**common, "SECTOR": "S1", "TRANSACTION": "EMP", "ACTIVITY": "_T", "UNIT_MEASURE": "PS"}, args.refresh),
        lambda: oecd_source("03_oecd_population", "DSD_NAMAIN10@DF_TABLE3_POP_EMPNC", "2.0", {**common, "SECTOR": "S1", "TRANSACTION": "POP", "UNIT_MEASURE": "PS"}, args.refresh),
        lambda: oecd_source("03_oecd_fiscal", "DSD_NASEC10@DF_TABLE12_EXP", "1.1", {**common, "SECTOR": "S13", "ACCOUNTING_ENTRY": "D", "TRANSACTION": "D1+P2+D632+OTE+D41", "UNIT_MEASURE": "XDC", "PRICE_BASE": "V"}, args.refresh),
        lambda: oecd_source("03_oecd_pib", "DSD_NAMAIN10@DF_TABLE1_EXPENDITURE", "2.0", {**common, "SECTOR": "S1", "TRANSACTION": "B1GQ", "UNIT_MEASURE": "XDC", "PRICE_BASE": "V"}, args.refresh),
        lambda: euro_source("03_eurostat_fiscal", "gov_10a_main", {"geo": EURO_GEOGRAPHIES, "sector": ["S13"], "unit": ["PC_GDP"], "na_item": ["D1PAY", "TE", "P2", "D632PAY"]}, args.refresh),
        lambda: euro_source("03_eurostat_cofog", "gov_10a_exp", {"geo": EURO_GEOGRAPHIES, "sector": ["S13"], "unit": ["PC_GDP"], "cofog99": [f"GF{i:02d}" for i in range(1, 11)], "na_item": ["TE", "D1", "P2", "D4"]}, args.refresh),
        lambda: euro_source("03_eurostat_emp", "gov_10a_emp", {"geo": EURO_GEOGRAPHIES, "sector": ["S13"], "unit": ["THS_PER"], "na_item": ["EMP_DC"]}, args.refresh),
    ]
    sources = []
    with ThreadPoolExecutor(max_workers=3) as pool:
        futures = [pool.submit(job) for job in jobs]
        for future in futures:
            source = future.result()
            sources.append(source)
            print(source["id"], len(source["rows"]), "observations", flush=True)
    lookup = {s["id"]: s for s in sources}
    gov, emp, pop, fiscal, gdp = [lookup["03_oecd_"+x] for x in ("emp_gov", "emp_total", "population", "fiscal", "pib")]
    sector_emp = lookup["03_oecd_emp_sector"]
    output, reconciliation = [], []
    for country in COUNTRIES:
        for year in YEARS:
            filters = {"REF_AREA": country, "TIME_PERIOD": year}
            numerator = unique(gov["rows"], **filters, SECTOR="S13", TRANSACTION="EMP")
            other_numerator = unique(sector_emp["rows"], **filters, SECTOR="S13", TRANSACTION="EMP")
            denominator = unique(emp["rows"], **filters, SECTOR="S1", TRANSACTION="EMP")
            population = unique(pop["rows"], **filters, TRANSACTION="POP")
            n = oecd_value(numerator) if numerator else None
            n2 = oecd_value(other_numerator) if other_numerator else None
            d = oecd_value(denominator) if denominator else None
            p = oecd_value(population) if population else None
            flags = ";".join(sorted({code(r.get("OBS_STATUS")) for r in (numerator, denominator, population) if r and r.get("OBS_STATUS")}))
            discrepancy_note = f"Écart avec table3 secteurs : {n-n2:g} personnes ; cause non résolue" if n is not None and n2 is not None and not math.isclose(n, n2, abs_tol=1) else ""
            if discrepancy_note:
                flags += ";ECART_TABLE3_NON_RESOLU"
            for indicator, source, value, item in (("emploi_s13_personnes", gov, n, numerator), ("emploi_total_interieur_personnes", emp, d, denominator), ("population_personnes", pop, p, population)):
                output.append(row_output(country, year, indicator, source, value, code(item.get("OBS_STATUS")) if item else "",
                                         note=discrepancy_note if indicator=="emploi_s13_personnes" else ""))
            output.append(row_output(country, year, "emploi_s13_pct_emploi", gov, n/d*100 if n is not None and d else None, flags,
                                     note=discrepancy_note,
                                     formula="100 × emploi S13 en personnes / emploi S1 intérieur en personnes", inputs=(emp,)))
            output.append(row_output(country, year, "emploi_s13_pour_1000_habitants", gov, n/p*1000 if n is not None and p else None, flags,
                                     note=discrepancy_note,
                                     formula="1000 × emploi S13 en personnes / population annuelle en personnes", inputs=(pop,)))
            if n is not None and n2 is not None:
                reconciliation.append(dict(pays_iso3=country, annee=year, emploi_table12=n, emploi_table3=n2, ecart=n-n2,
                                           source_url=gov["meta"]["url"]+" ; "+sector_emp["meta"]["url"],
                                           source_localisation="03_oecd_emp_gov.csv ; 03_oecd_emp_sector.csv", statut="[FAIT]"))
            pib_row = unique(gdp["rows"], **filters, TRANSACTION="B1GQ")
            pib = oecd_value(pib_row) if pib_row else None
            ratios = {}
            fiscal_flags = set()
            for trans, indicator in (("D1", "D1_pct_pib"), ("OTE", "TE_pct_pib"), ("P2", "P2_pct_pib"), ("D632", "D632_pct_pib")):
                raw = unique(fiscal["rows"], **filters, TRANSACTION=trans)
                amount = oecd_value(raw) if raw else None
                currency = code(raw.get("CURRENCY")) if raw else ""
                same_currency = raw and pib_row and currency not in ("", "_Z") and currency == code(pib_row.get("CURRENCY"))
                ratio = amount / pib * 100 if amount is not None and pib and same_currency else None
                flags = ";".join(sorted({code(r.get("OBS_STATUS")) for r in (raw, pib_row) if r and r.get("OBS_STATUS")}))
                fiscal_flags.update(x for x in flags.split(";") if x)
                ratios[trans] = ratio
                output.append(row_output(country, year, indicator, fiscal, ratio, flags,
                                         note="Même monnaie nationale vérifiée ; valeurs et multiplicateurs conservés dans les bruts" if ratio is not None else "Numérateur, PIB ou concordance de monnaie absent",
                                         formula="100 × montant S13 en monnaie nationale / PIB courant même monnaie et année", inputs=(gdp,)))
            for components, indicator in ((["P2", "D632"], "P2_D632_pct_pib"), (["D1", "P2", "D632"], "D1_P2_D632_pct_pib")):
                value = sum(ratios[c] for c in components) if all(ratios[c] is not None for c in components) else None
                output.append(row_output(country, year, indicator, fiscal, value, ";".join(sorted(fiscal_flags)), formula=" + ".join(c+"/%PIB" for c in components), inputs=(gdp,)))
    for name, trans_map in (("fiscal", {"D1PAY":"D1_pct_pib", "TE":"TE_pct_pib", "P2":"P2_pct_pib", "D632PAY":"D632_pct_pib"}),):
        source = lookup["03_eurostat_"+name]
        for geo in EURO_GEOGRAPHIES:
            country = ISO2_TO_ISO3.get(geo, geo)
            for year in YEARS:
                for transaction, indicator in trans_map.items():
                    row = unique(source["rows"], geo=geo, time=year, na_item=transaction)
                    output.append(row_output(country, year, indicator, source, row["value"] if row else None, row.get("flag", "") if row else ""))
    source = lookup["03_eurostat_emp"]
    for geo in EURO_GEOGRAPHIES:
        country = ISO2_TO_ISO3.get(geo, geo)
        for year in YEARS:
            row = unique(source["rows"], geo=geo, time=year)
            output.append(row_output(country, year, "emploi_s13_personnes", source, row["value"]*1000 if row else None,
                                     row.get("flag", "") if row else "", formula="milliers de personnes × 1000"))
    source = lookup["03_eurostat_cofog"]
    for function in (f"GF{i:02d}" for i in range(1, 11)):
        for transaction in ("TE", "D1", "P2", "D4"):
            indicator = f"{function}_{transaction}_pct_pib"
            SCHEMAS[indicator] = ("% PIB", f"Dépense {transaction} S13, fonction COFOG {function} ; ni catégorie de métier ni effectif administratif")
            for geo in EURO_GEOGRAPHIES:
                country = ISO2_TO_ISO3.get(geo, geo)
                for year in YEARS:
                    row = unique(source["rows"], geo=geo, time=year, cofog99=function, na_item=transaction)
                    output.append(row_output(country, year, indicator, source, row["value"] if row else None, row.get("flag", "") if row else ""))
    for geo in EURO_GEOGRAPHIES:
        country = ISO2_TO_ISO3.get(geo, geo)
        for year in YEARS:
            output.append(row_output(country, year, "GF01_hors_interets_pct_pib", source, None, "D41_FONCTIONNEL_NON_TROUVE",
                                     "D4 n'est pas exclusivement D41 ; ne pas soustraire GF0107 intégral"))
    publication = panorama_source(args.refresh)
    sources.append(publication)
    panorama_checks = []
    for item in publication["rows"]:
        country, year = item["country"], item["year"]
        note = "Observation publiée dans la figure13.1 ; source et millésime distincts des API actuelles"
        flag = "ANNEE_REELLE_"+str(year)+"_COLONNE_2023" if year != 2023 else ""
        if country == "JPN":
            note += " ; le fichier publié précise : fonds de sécurité sociale exclus"
        if country == "AUS":
            note += " ; emploi total fondé sur une estimation OCDE"
        if country in {"AUS", "ISL", "JPN", "KOR", "TUR", "USA"}:
            note += " ; données ILOSTAT reprises par l'OCDE"
        if country == "CHL":
            note += " ; effectif public communiqué par les autorités nationales"
        if country == "OECD":
            note += " ; moyenne publiée : certains pays sont exclus, voir notes des cellules A24:A26"
        row = row_output(country, year, "emploi_s13_pct_emploi", publication, item["value"], flag, note)
        row["source_localisation"] = str(publication["path"].relative_to(ROOT)).replace("\\", "/") + " ; " + item["cell"] + " ; notes A22:A26"
        output.append(row)
        current = next((r for r in output if r["pays_iso3"] == country and r["annee"] == year and r["indicateur"] == "emploi_s13_pct_emploi" and r["source_id"] == gov["id"] and r["valeur"] != ""), None)
        if current:
            panorama_checks.append(dict(pays_iso3=country, annee=year, valeur_publication_2025=item["value"], valeur_api_2026=current["valeur"],
                                        ecart_points=float(current["valeur"])-item["value"],
                                        source_url=publication["meta"]["url"]+" ; "+current["source_url"],
                                        source_localisation=item["cell"]+" ; "+current["source_localisation"], statut="[ESTIMATION]",
                                        note="Différence de sources/vintages/périmètres à examiner ; pas une évolution réelle entre deux années"))
    write_csv(OUT / "03_base_comparative.csv", output)
    write_csv(OUT / "03_rapprochement_panorama2025.csv", panorama_checks)
    finance_checks = financial_reconciliation(output)
    keys = [(r["source_id"], r["pays_iso3"], r["annee"], r["indicateur"]) for r in output]
    if len(keys) != len(set(keys)):
        raise ValueError("Clé d'observation dupliquée dans la base finale")
    if any(r["valeur"] == "" and r["statut"] != "[NON TROUVÉ]" for r in output):
        raise ValueError("Valeur manquante sans statut explicite")
    if reconciliation:
        write_csv(OUT / "03_rapprochement_emplois_s13.csv", reconciliation)
    write_csv(OUT / "03_dictionnaire_indicateurs.csv", [dict(indicateur=k, unite=v[0], definition=v[1], annee="2023-2024", statut="[FAIT]", source_url="Voir source_url de la base", source_localisation="Définitions opérationnelles du script 03_collecte_harmonisee.py") for k,v in SCHEMAS.items()])
    variable_definitions = {
        "pays_iso3": "Code ISO3 ; EU27_2020 et OECD désignent des agrégats, pas des pays",
        "pays": "Libellé français du pays ou de l'agrégat",
        "indicateur": "Identifiant stable de l'indicateur, voir 03_dictionnaire_indicateurs.csv",
        "valeur": "Valeur numérique ; vide signifie indisponible, jamais zéro imputé",
        "unite": "Unité après prise en compte du multiplicateur de la source",
        "annee": "Année réelle de référence ; les dérogations StatLink 2021/2022 sont corrigées explicitement",
        "champ": "Définition opérationnelle et avertissement sur le champ",
        "secteur": "S13 administrations publiques ou S1 économie/population totale",
        "origine": "Institution productrice/diffuseuse ; ne garantit pas une source nationale indépendante",
        "source_id": "Clé du registre03 ; indispensable pour distinguer sources, reprises et éditions",
        "source_url": "URL exacte des réponses utilisées ; plusieurs URL séparées par point-virgule pour un calcul",
        "edition_source": "Publication datée ou collecte API ; ne pas fusionner silencieusement les millésimes",
        "source_localisation": "Chemin local, filtre pays/année ou cellule Excel et notes",
        "flag": "Code source conservé ; A normal, P/p provisoire ; autres avertissements explicites du calcul",
        "statut": "FAIT lu dans la source, ESTIMATION calcul reproductible, NON TROUVÉ absence ou formule impossible",
        "formule": "Formule appliquée ; vide pour une valeur publiée directement",
        "note": "Limites, divergences ou exceptions liées à cette observation",
        "date_consultation": "Date de la mission au format ISO ; actualisation future : réviser explicitement DATE dans le script",
    }
    write_csv(OUT / "03_dictionnaire_variables.csv", [dict(variable=k, definition=v) for k,v in variable_definitions.items()])
    registry = []
    for source in sources:
        metadata = [(source["id"], source["path"], source["meta"], "Réponse de données primaire intégralement lue ; filtres exacts dans l'URL")]
        if "structure_path" in source:
            metadata.append((source["id"]+"_STRUCTURE", source["structure_path"], source["structure_meta"], "Structure primaire lue pour déterminer l'ordre des dimensions"))
        for sid, path, meta, note in metadata:
            registry.append(dict(id=sid, titre=path.name, institution=source["institution"], date_publication=source.get("publication", ""),
                                 url=meta["url"], date_consultation=meta["date_consultation"], acces="integral", sha256_16=meta["sha256"][:16],
                                 notes=note+" ; "+str(path.relative_to(ROOT)).replace("\\", "/")+
                                 (" ; mise à jour API : "+source["updated"] if source.get("updated") else " ; date de mise à jour API non renseignée")))
    write_csv(ROOT / "sources/registre_03.csv", registry)
    print("Base :",len(output),"lignes ; valeurs disponibles",sum(r["valeur"]!="" for r in output), flush=True)
    print("Rapprochements S13 :",len(reconciliation),"; écarts",sum(abs(r["ecart"])>1 for r in reconciliation), flush=True)
    print("Rapprochements financiers :", len(finance_checks), "; écarts >0,15 point", sum(abs(r["ecart_points"])>0.15 for r in finance_checks), flush=True)


if __name__ == "__main__":
    main()
