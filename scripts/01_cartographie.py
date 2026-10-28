#!/usr/bin/env python3
"""Chantier 01 : extraction déterministe des XLSX hérités, sans connexion réseau.

Les cellules, unités originales, multiplicateurs et versions restent dans chaque ligne.
Ne pas sommer les lignes portant niveau=total/sous_total/detail entre elles.
"""
from __future__ import annotations
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/01"
OUT = ROOT / "data/france"
OUT.mkdir(parents=True, exist_ok=True)
DG = "https://www.fonction-publique.gouv.fr/files/files/publications/vue-rapport-annuel/2026/"
SOURCES = {
    "insee_IP2094_donnees.xlsx": ("https://www.insee.fr/fr/statistiques/fichier/8732435/IP2094.xlsx", "2026-02-10"),
    "insee_2493501_marc-empl-eff-fonc-pub_2011-2024.xlsx": ("https://www.insee.fr/fr/statistiques/fichier/2493501/marc-empl-eff-fonc-pub.xlsx", "2026-02-10"),
    "insee_2381404_depenses_apu_2025.xlsx": ("https://www.insee.fr/fr/statistiques/fichier/2381404/econ-fin-pub-dep-adm-pub-secteur.xlsx", "2026-08-28"),
    "ra_2026_vue_emploi_effectifs.xlsx": (DG + "ra_2026_effectifs.zip", "2026-10-15"),
    "RA_2026_effectifs_donnees_comp_FP.xlsx": (DG + "ra_2026_effectifs.zip", "2026-10-15"),
    "ra_2026_vue_emploi_caracteristiques_postes.xlsx": (DG + "ra_2026_postes.zip", "2026-10-15"),
    "ra_2026_remu_niveaux_donnees_comp_depenses.xlsx": (DG + "ra_2026_remu_niveaux.zip", "2026-10-15"),
    "ra_2026_remu_niveaux_donnees_comp.xlsx": (DG + "ra_2026_remu_niveaux.zip", "2026-10-15"),
    "ra_2026_vue_temps_organisation_temps_travail.xlsx": (DG + "ra_2026_tps_organisation_tps_travail.zip", "2026-10-15"),
}
BOOKS = {}
PATHS = {p.name: p for p in RAW.rglob("*.xlsx")}
FIELDS = ["id", "indicateur", "versant", "categorie", "annee", "valeur", "unite", "champ", "niveau", "version",
          "source_url", "source_localisation", "date_publication", "statut", "fichier_brut", "valeur_brute", "multiplicateur", "notes"]


def book(name):
    if name not in BOOKS:
        BOOKS[name] = load_workbook(PATHS[name], read_only=False, data_only=True)
    return BOOKS[name]


def val(name, sheet, cell):
    x = book(name)[sheet][cell].value
    if not isinstance(x, (int, float)):
        raise ValueError(f"Cellule non numérique : {name}!{sheet}!{cell}={x!r}")
    return x


def record(name, sheet, cell, ident, label, year, unit, versant="FP", category="", multiplier=1,
           champ="France hors Mayotte, postes principaux au 31 décembre, hors contrats aidés",
           level="detail", note="", version=""):
    raw = val(name, sheet, cell)
    url, date = SOURCES[name]
    return dict(zip(FIELDS, [ident, label, versant, category, year, round(raw * multiplier, 8), unit, champ, level,
                           version or ("DGAFP RA 2026" if "insee" not in name else "Insee publication 2026"),
                           url, f"{name} / {sheet}!{cell}", date, "[FAIT]", str(PATHS[name].relative_to(ROOT)), raw, multiplier, note]))


def write(name, rows, fields=FIELDS):
    with (OUT / name).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main():
    counts = {}
    rows = []
    source = "insee_IP2094_donnees.xlsx"
    for year, cols in [(2023, ("B", "C")), (2024, ("D", "E"))]:
        for inc, col in zip((True, False), cols):
            for row, versant in zip(range(6, 10), ("FPE", "FPT", "FPH", "FP")):
                field = "y compris contrats aidés" if inc else "hors contrats aidés"
                for offset, label, unit, scope in [(0, "Effectif au 31 décembre", "personnes physiques", "postes principaux au 31 décembre"),
                                                  (5, "Volume de travail annuel", "EQTP annuels", "tout poste actif dans l'année")]:
                    rows.append(record(source, "Figure 1", f"{col}{row+offset}", f"IP2094_{year}_{versant}_{inc}_{offset}",
                                       label, year, unit, versant, field, 1000, f"France hors Mayotte, {scope}, {field}",
                                       level="total" if versant == "FP" else "detail",
                                       note="L'EQTP annuel de tous les postes n'est pas l'ETP instantané du stock au 31 décembre." if offset else ""))
    source = "RA_2026_effectifs_donnees_comp_FP.xlsx"
    for year, col in [(2023, "O"), (2024, "P")]:
        for row, versant in [(15,"FPE"),(16,"FPT"),(17,"FPH"),(18,"FP"),(19,"FPE"),(20,"FPT"),(21,"FPH"),(22,"FP")]:
            aid = "y compris contrats aidés" if row <= 18 else "hors contrats aidés"
            rows.append(record(source, "Figure 4", f"{col}{row}", f"DGAFP_{year}_{versant}_{row}", "Effectif au 31 décembre",
                               year, "personnes physiques", versant, aid, 1000,
                               f"France hors Mayotte, postes principaux au 31 décembre, {aid}", level="total" if versant=="FP" else "detail"))
    counts["01_effectifs_versions.csv"] = write("01_effectifs_versions.csv", rows)

    rows=[]
    source="ra_2026_vue_emploi_caracteristiques_postes.xlsx"
    for versant, mapping in [("FPE", [(6,"fonctionnaires"),(7,"contractuels"),(8,"militaires"),(9,"autres"),(10,"total")]),
                             ("FPT", [(11,"fonctionnaires"),(12,"contractuels"),(13,"autres"),(14,"total")]),
                             ("FPH", [(15,"fonctionnaires"),(16,"contractuels"),(17,"autres"),(18,"total")]),
                             ("FP", [(19,"fonctionnaires"),(20,"contractuels"),(21,"militaires"),(22,"autres"),(23,"total")])]:
        for row,statut in mapping:
            rows.append(record(source,"Figure 1",f"D{row}",f"STATUT_{versant}_{statut}","Effectif par statut",2024,"personnes physiques",versant,statut,1000,
                               level="total" if statut=="total" else "detail",
                               note="FPT : militaires regroupés avec les fonctionnaires et volontaires avec les contractuels pour secret statistique. Autres inclut médecins et enseignants privés sous contrat."))
    counts["01_effectifs_statuts.csv"]=write("01_effectifs_statuts.csv",rows)

    rows=[]
    source="ra_2026_vue_emploi_effectifs.xlsx"
    for sheet,start,end,versant in [("Figure 4",6,37,"FPE"),("Figure 6",6,17,"FPT"),("Figure 7",6,11,"FPH")]:
        parent=""
        for row in range(start,end+1):
            label=book(source)[sheet][f"B{row}"].value
            if not label:
                continue
            if not label.lstrip().startswith("dont"):
                parent=label.strip()
            level="detail_inclus" if label.lstrip().startswith("dont") else ("total" if "Ensemble" in label or "Total" in label else "sous_total")
            rows.append(record(source,sheet,f"C{row}",f"EMPLOYEUR_{versant}_{row}","Effectifs par employeur",2024,"personnes physiques",versant,
                               f"{parent} / {label.strip()}" if level=="detail_inclus" else label.strip(),1000,level=level,
                               note="Mission de l'employeur, pas métier de l'agent. Les sous-totaux et lignes 'dont' ne sont pas additifs."))
    counts["01_effectifs_employeurs.csv"]=write("01_effectifs_employeurs.csv",rows)

    rows=[]
    source="RA_2026_effectifs_donnees_comp_FP.xlsx"
    for row,versant in [(17,"FPE"),(20,"FPT"),(21,"FPH"),(22,"FP")]:
        for col in range(18,32):
            year=int(str(book(source)["Figure 5"].cell(14,col).value)[:4])
            rows.append(record(source,"Figure 5",f"{get_column_letter(col)}{row}",f"SERIE_{versant}_{year}","Effectif série rétropolée",year,"personnes physiques",versant,
                               "hors contrats aidés",1000,level="total" if versant=="FP" else "detail",
                               note="Série rétropolée publiée par DGAFP. Ne pas raccorder aux niveaux anciens non rétropolés. AESH : transferts EPA/ministères sans création équivalente d'emploi."))
    counts["01_effectifs_serie_2011_2024.csv"]=write("01_effectifs_serie_2011_2024.csv",rows)

    rows=[]
    for row in range(9,29):
        label=str(book(source)["Figure 3"][f"C{row}"].value).strip()
        level="total" if row in [20,25,28] else ("sous_total" if row in [16,19] else "detail")
        group="Organismes publics hors FP" if row<21 else ("Organismes privés à financement public prédominant" if row<26 else "Entreprises publiques marchandes")
        rows.append(record(source,"Figure 3",f"U{row}",f"HORS_FP_{row}",group,2024,"personnes physiques","hors FP",label,1000,
                           "France hors Mayotte, emplois principaux au 31 décembre, contrats aidés inclus",level=level,
                           note="Ne pas ajouter au stock S13 : champs partiellement imbriqués. Ne représente pas tous les prestataires privés."))
    rows.append(record(source,"Figure 2","L28","EMPLOI_S13","Effectif des administrations publiques au sens des comptes nationaux",2024,"personnes physiques","S13","ensemble",1000,
                       "France hors Mayotte, salariés à titre principal au 31 décembre, contrats aidés inclus",level="total",note="Champ S13 distinct de la fonction publique. Pas un effectif administratif."))
    counts["01_effectifs_hors_fp.csv"]=write("01_effectifs_hors_fp.csv",rows)

    rows=[]
    source="insee_2381404_depenses_apu_2025.xlsx"
    for row,measure in [(5,"Dépenses totales"),(7,"Consommations intermédiaires P2"),(8,"Rémunération des salariés D1")]:
        for col,sector in [("B","APUC"),("C","Etat_dont_APUC"),("D","APUL"),("E","ASSO"),("F","S13")]:
            rows.append(record(source,"Figure 1",f"{col}{row}",f"CN2025_{measure}_{sector}",measure,2025,"MdEUR courants",sector,level="total" if sector=="S13" else "detail",
                               champ="France, comptes nationaux base 2020, données provisoires",note="D1 inclut cotisations employeur effectives et imputées ; ni coût marginal d'entrée ni masse salariale hors pensions. Etat inclus dans APUC."))
    source="ra_2026_remu_niveaux_donnees_comp_depenses.xlsx"
    for year,col in [(2024,"AV"),(2025,"AW")]:
        for row,sector in [(5,"APUC"),(6,"Etat_dont_APUC"),(7,"APUL"),(9,"ASSO"),(11,"S13")]:
            rows.append(record(source,"Figure 1",f"{col}{row}",f"DG_D1_{year}_{sector}","Rémunération des salariés D1",year,"MdEUR courants",sector,
                               champ="France, comptes nationaux base 2020",level="total" if sector=="S13" else "detail",note="Source initiale Insee : contrôle de transcription, pas source statistique indépendante."))
    for year,col in [(2024,"S"),(2025,"T")]:
        for row in [4,13,14,15,16,17,19,20,21,22,23,24,25,35,37,38]:
            label=book(source)["Figure 2"][f"B{row}"].value
            rows.append(record(source,"Figure 2",f"{col}{row}",f"T2_{year}_{row}",label,year,"MdEUR courants","Etat_budget_general",champ="Budget général de l'Etat, budget d'exécution",
                               level="total" if row==35 else "detail",note="Lignes imbriquées : ne pas sommer total, cotisations et sous-postes. Pensions versées (ligne37) distinctes des contributions CAS (ligne14)."))
    counts["01_depenses_personnel.csv"]=write("01_depenses_personnel.csv",rows)

    rows=[]
    source="ra_2026_remu_niveaux_donnees_comp.xlsx"
    for versant,all_row,hca_row,young_row in [("FPE",4,11,15),("FPT",21,28,32),("FPH",38,45,49)]:
        for row,label in [(all_row,"Salaire brut moyen, champ publié"),(hca_row,"Salaire brut moyen hors contrats aidés"),(young_row,"Salaire brut moyen des moins de 30 ans")]:
            rows.append(record(source,"Figure 1",f"B{row}",f"BRUT_{versant}_{row}",label,2024,"EUR par mois et EQTP mensualisé",versant,
                               champ="France hors Mayotte, civils, hors assistants maternels/familiaux, apprentis, internes/externes ; source SIASP",note="Salaire brut sans charges employeur. Moins de 30 ans n'est pas un coût d'entrée. Tous métiers, pas uniquement administratifs."))
    source="ra_2026_vue_temps_organisation_temps_travail.xlsx"
    for col,versant in [("C","FPE"),("D","FPT"),("E","FPH"),("F","FP"),("G","prive")]:
        rows.append(record(source,"Figure 1",f"{col}6",f"HEURES_{versant}","Durée annuelle effective déclarée",2025,"heures par an",versant,
                           champ="Enquête Emploi, temps complet, hors enseignants, apprentis, contrats aidés ; France hors Mayotte, logement ordinaire",note="Déclaratif, structure des métiers différente du privé. Ni durée légale ni preuve de postes supprimables."))
    counts["01_salaires_temps.csv"]=write("01_salaires_temps.csv",rows)

    # Rapprochement arithmétique des versions. Une concordance n'explique pas causalement une révision.
    checks=[]
    def control(label,observed,expected,tolerance=0.01):
        checks.append(dict(controle=label,observe=round(observed,8),attendu=round(expected,8),ecart=round(observed-expected,8),
                           tolerance=tolerance,resultat="OK" if abs(observed-expected)<=tolerance else "ECART_A_DOCUMENTER"))
    source="ra_2026_vue_emploi_effectifs.xlsx"
    control("DGAFP somme des versants",sum(val(source,"Figure 2",f"C{r}") for r in [6,7,8]),val(source,"Figure 2","C9"))
    for versant,rows_,total in [("FPE",[6,7,8,9],10),("FPT",[11,12,13],14),("FPH",[15,16,17],18),("FP",[19,20,21,22],23)]:
        f="ra_2026_vue_emploi_caracteristiques_postes.xlsx"
        control(f"Statuts {versant}",sum(val(f,"Figure 1",f"D{r}") for r in rows_),val(f,"Figure 1",f"D{total}"))
    control("FPE somme rattachements ministériels",sum(val(source,"Figure 4",f"C{r}") for r in [6,9,12,15,18,21,24,27,30,33,34]),val(source,"Figure 2","C6"))
    control("FPT somme groupes employeurs",sum(val(source,"Figure 6",f"C{r}") for r in [6,9,12,15,16]),val(source,"Figure 2","C7"))
    control("FPH somme groupes employeurs",sum(val(source,"Figure 7",f"C{r}") for r in [6,9,10]),val(source,"Figure 2","C8"))
    f="insee_2381404_depenses_apu_2025.xlsx"
    control("D1 2025 somme APUC APUL ASSO",sum(val(f,"Figure 1",f"{c}8") for c in "BDE"),val(f,"Figure 1","F8"))
    control("D1 2025 Insee vs DGAFP",val(f,"Figure 1","F8"),val("ra_2026_remu_niveaux_donnees_comp_depenses.xlsx","Figure 1","AW11"))
    counts["01_controles.csv"]=write("01_controles.csv",checks,list(checks[0]))
    manifest=[dict(fichier=str(PATHS[n].relative_to(ROOT)),sha256=hashlib.sha256(PATHS[n].read_bytes()).hexdigest(),source_url=u,date_publication=d)
              for n,(u,d) in SOURCES.items()]
    counts["01_manifest_fichiers.csv"]=write("01_manifest_fichiers.csv",manifest,list(manifest[0]))
    for b in BOOKS.values():
        b.close()
    print(json.dumps(counts,ensure_ascii=False))
    if any(c["resultat"]!="OK" for c in checks):
        raise SystemExit("Écart de cohérence à examiner dans 01_controles.csv")


if __name__ == "__main__":
    main()
