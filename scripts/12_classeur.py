#!/usr/bin/env python3
"""Écrit modele/modele_financier.xlsx : version tableur, modifiable, du catalogue.

Le classeur calcule un **régime permanent simplifié** par formules : cibles successives
dans chaque vivier, plafond, capacité de départs sur l'horizon, économie au coût d'entrée
vieilli d'une demi-période. Il ne remplace pas modele/modele.py (calcul annuel par cohorte) ;
la feuille Synthèse affiche les deux résultats et leur écart. Cellules jaunes = hypothèses.
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SC = ["prudent", "intermediaire", "ambitieux", "radical"]
LAB = {"prudent": "Prudent", "intermediaire": "Intermédiaire", "ambitieux": "Ambitieux", "radical": "Radical"}
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="DCE6EF")
BOLD = Font(bold=True)
SEUILS = [50_000, 100_000, 200_000, 500_000, 800_000, 1_000_000, 1_500_000]


def read(path):
    with open(ROOT / path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def header(ws, row, labels):
    for j, label in enumerate(labels, 1):
        c = ws.cell(row, j, label)
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")


REFS = {}


def name(wb, label, ref):
    """Références absolues plutôt que plages nommées : lisibles par tous les tableurs et moteurs de recalcul."""
    REFS[label] = ref


def f(formula):
    for label in sorted(REFS, key=len, reverse=True):
        formula = formula.replace(label, REFS[label])
    return formula


def widths(ws, values):
    for j, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def main():
    hyp = json.loads((ROOT / "modele/hypotheses.json").read_text(encoding="utf-8"))
    viviers = read("data/reformes/registre_viviers.csv")
    reformes = sorted(read("data/reformes/catalogue_reformes.csv"), key=lambda r: (r["vivier_id"], float(r["ordre"])))
    synth = {r["scenario"]: r for r in read("modele/resultats/synthese_scenarios.csv")}
    seuils_csv = read("data/reformes/12_seuils.csv")
    sens = read("data/reformes/12_sensibilites.csv")
    wb = Workbook()

    ws = wb.active
    ws.title = "Lisez-moi"
    notes = [
        "Modèle financier du catalogue de réformes — version tableur (30 octobre 2026)",
        "",
        "Les cellules jaunes sont des hypothèses modifiables (feuilles Hypothèses, Viviers, Réformes). Tout le reste est calculé.",
        "Le classeur calcule un régime permanent simplifié à l'horizon choisi : il ne simule pas les cohortes année par année.",
        "Le moteur de référence est modele/modele.py ; la feuille Synthèse compare les deux calculs et affiche leur écart.",
        "Convention : économies = écart à une référence qui remplace tous les départs, en M€ 2026 constants.",
        "Ordre d'application : dans chaque vivier, chaque réforme porte sur le reliquat laissé par les précédentes (lignes triées par vivier puis ordre).",
        "Capacité : le non-remplacement ne peut dépasser les départs mobilisables du vivier, B × (1 − (1 − r × multiplicateur)^horizon).",
        "Coût d'un poste non remplacé : MIN(coût moyen ; coût d'entrée × (1 + GVT)^(horizon / 2)) × multiplicateur de coût.",
        "Sources et justification de chaque paramètre : recherche/11_catalogue_reformes.md et recherche/12_scenarios_seuils.md.",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(i, 1, t).font = BOLD if i == 1 else Font()
    ws.column_dimensions["A"].width = 130

    # --- Hypothèses
    h = wb.create_sheet("Hypotheses")
    header(h, 1, ["Paramètre", "Valeur", "Source / justification"])
    params = [("Horizon (années)", 10, "2027-2036, comme le moteur", "horizon"),
              ("GVT individuel réel", hyp["gvt_individuel_reel"], hyp["_gvt_source"], "gvt"),
              ("Conservation APU (lecture consolidée)", hyp["taux_conservation_consolide"], hyp["_conservation_source"], "conservation"),
              ("Multiplicateur des taux de réduction", 1, "1 = catalogue ; sensibilité testée 0,5-1,5", "mult_taux"),
              ("Multiplicateur des départs mobilisables", 1, "1 = retraites ; 1,605 = toutes sorties de fonctionnaires", "mult_departs"),
              ("Multiplicateur des coûts par ETP", 1, "sensibilité testée 0,8-1,2", "mult_cout"),
              ("Multiplicateur des investissements", 1, "sensibilité testée 0,5-2", "mult_invest")]
    for i, (label, value, src, nm) in enumerate(params, 2):
        h.cell(i, 1, label)
        c = h.cell(i, 2, value)
        c.fill = INPUT
        h.cell(i, 3, src)
        name(wb, nm, f"Hypotheses!$B${i}")
    r0 = len(params) + 3
    header(h, r0, ["Versant", "Coût d'entrée hors pension (k€)", "Coût moyen hors pension (k€)", "Source"])
    for i, (versant, c) in enumerate(hyp["couts_k_eur"].items(), r0 + 1):
        h.cell(i, 1, versant)
        for j, k in ((2, "cout_entree_hors_pension"), (3, "cout_moyen_hors_pension")):
            h.cell(i, j, c[k]).fill = INPUT
        h.cell(i, 4, c["_source"])
    couts = f"Hypotheses!$A${r0 + 1}:$C${r0 + len(hyp['couts_k_eur'])}"
    widths(h, [42, 16, 16, 120])

    # --- Viviers
    v = wb.create_sheet("Viviers")
    cols = ["vivier_id", "versant", "etp_base", "taux_departs_annuel", *[f"plafond_reduction_{s}" for s in SC], "Capacité de départs sur l'horizon (ETP)", "description", "sources"]
    header(v, 1, cols)
    for i, row in enumerate(viviers, 2):
        v.cell(i, 1, row["vivier_id"])
        v.cell(i, 2, row["versant"])
        for j, k in enumerate(["etp_base", "taux_departs_annuel", *[f"plafond_reduction_{s}" for s in SC]], 3):
            v.cell(i, j, float(row[k])).fill = INPUT
        v.cell(i, 9, f(f"=C{i}*(1-(1-MIN(1,D{i}*mult_departs))^horizon)"))
        v.cell(i, 10, row["description"])
        v.cell(i, 11, row["sources"])
    nv = len(viviers) + 1
    v.cell(nv + 1, 1, "Total").font = BOLD
    v.cell(nv + 1, 3, f"=SUM(C2:C{nv})")
    v.cell(nv + 1, 9, f"=SUM(I2:I{nv})")
    viv_tab = f"Viviers!$A$2:$I${nv}"
    widths(v, [14, 11, 12, 12, 11, 11, 11, 11, 18, 50, 80])

    # --- Réformes (entrées)
    rf = wb.create_sheet("Reformes")
    rcols = ["id", "intitule", "vivier_id", "ordre", "part_non_remplacement", "part_redeploiement", "cout_recurrent_proportionnel",
             *[f"taux_{s}" for s in SC], *[f"invest_meur_{s}" for s in SC], *[f"cout_recurrent_meur_{s}" for s in SC],
             *[f"eco_fonct_meur_{s}" for s in SC], *[f"effet_recettes_meur_{s}" for s in SC], "confiance", "fiche"]
    header(rf, 1, rcols)
    for i, row in enumerate(reformes, 2):
        for j, k in enumerate(rcols, 1):
            val = row[k]
            if k not in ("id", "intitule", "vivier_id", "confiance", "fiche"):
                val = float(val)
            c = rf.cell(i, j, val)
            if k.startswith(("taux_", "invest_", "cout_recurrent_meur", "eco_fonct", "effet_recettes")):
                c.fill = INPUT
    nr = len(reformes) + 1
    col = {k: get_column_letter(j) for j, k in enumerate(rcols, 1)}
    widths(rf, [6, 60, 14] + [10] * (len(rcols) - 3))

    # --- Calcul par scénario
    for s in SC:
        c = wb.create_sheet(f"Calcul_{s}")
        ccols = ["id", "vivier", "ordre", "base ETP", "plafond", "taux effectif", "cumul cibles précédentes", "cible ETP",
                 "non-remplacement demandé", "capacité départs vivier", "cumul NR précédents", "NR réalisé", "redéployés",
                 "coût unitaire (k€)", "économie salariale RP (M€)", "éco. fonctionnement (M€)", "coût récurrent (M€)",
                 "recettes (M€)", "net RP (M€)", "consolidé APU RP (M€)", "investissement (M€)"]
        header(c, 1, ccols)
        vi = {k: j for j, k in enumerate(["vivier_id", "versant", "etp_base", "taux_departs_annuel", *[f"plafond_reduction_{x}" for x in SC], "cap"], 1)}
        for i in range(2, nr + 1):
            R = f"Reformes!"
            c.cell(i, 1, f"={R}A{i}")
            c.cell(i, 2, f"={R}C{i}")
            c.cell(i, 3, f"={R}D{i}")
            c.cell(i, 4, f"=VLOOKUP(B{i},{viv_tab},{vi['etp_base']},FALSE)")
            c.cell(i, 5, f"=VLOOKUP(B{i},{viv_tab},{vi['plafond_reduction_' + s]},FALSE)")
            c.cell(i, 6, f(f"=MIN(1,{R}{col['taux_' + s]}{i}*mult_taux)"))
            c.cell(i, 7, "=0" if i == 2 else f"=IF(B{i}=B{i - 1},G{i - 1}+H{i - 1},0)")
            c.cell(i, 8, f"=MAX(0,MIN(F{i}*(D{i}-G{i}),E{i}*D{i}-G{i}))")
            c.cell(i, 9, f"=H{i}*{R}{col['part_non_remplacement']}{i}")
            c.cell(i, 10, f"=VLOOKUP(B{i},{viv_tab},{vi['cap']},FALSE)")
            c.cell(i, 11, "=0" if i == 2 else f"=IF(B{i}=B{i - 1},K{i - 1}+L{i - 1},0)")
            c.cell(i, 12, f"=MIN(I{i},MAX(0,J{i}-K{i}))")
            c.cell(i, 13, f"=H{i}*{R}{col['part_redeploiement']}{i}")
            versant = f"VLOOKUP(B{i},{viv_tab},{vi['versant']},FALSE)"
            c.cell(i, 14, f(f"=MIN(VLOOKUP({versant},{couts},3,FALSE),VLOOKUP({versant},{couts},2,FALSE)*(1+gvt)^(horizon/2))*mult_cout"))
            c.cell(i, 15, f"=L{i}*N{i}/1000")
            c.cell(i, 16, f"={R}{col['eco_fonct_meur_' + s]}{i}")
            c.cell(i, 17, f"=IF({R}{col['cout_recurrent_proportionnel']}{i}=1,{R}{col['cout_recurrent_meur_' + s]}{i}*IF(I{i}>0,L{i}/I{i},0),{R}{col['cout_recurrent_meur_' + s]}{i})")
            c.cell(i, 18, f"={R}{col['effet_recettes_meur_' + s]}{i}")
            c.cell(i, 19, f"=O{i}+P{i}-Q{i}+R{i}")
            c.cell(i, 20, f(f"=O{i}*conservation+P{i}-Q{i}+R{i}"))
            c.cell(i, 21, f(f"={R}{col['invest_meur_' + s]}{i}*mult_invest"))
        t = nr + 1
        c.cell(t, 1, "Total").font = BOLD
        for j in (8, 12, 13, 15, 16, 17, 18, 19, 20, 21):
            L = get_column_letter(j)
            c.cell(t, j, f"=SUM({L}2:{L}{nr})").font = BOLD
        widths(c, [6, 14, 7] + [13] * 18)
        name(wb, f"postes_{s}", f"Calcul_{s}!$L${t}")
        name(wb, f"net_{s}", f"Calcul_{s}!$S${t}")
        name(wb, f"conso_{s}", f"Calcul_{s}!$T${t}")
        name(wb, f"invest_{s}", f"Calcul_{s}!$U${t}")

    # --- Synthèse
    sy = wb.create_sheet("Synthese", 1)
    header(sy, 1, ["Indicateur", *[LAB[s] for s in SC]])
    lines = [("Postes en moins à l'horizon — tableur (régime permanent)", "postes_{}"),
             ("Postes en moins en 2036 — moteur", None, "postes_en_moins_10ans"),
             ("Économie nette annuelle à l'horizon (M€) — tableur", "net_{}"),
             ("Économie nette 2036 (M€) — moteur", None, "eco_nette_2036_meur"),
             ("Économie consolidée APU à l'horizon (M€) — tableur", "conso_{}"),
             ("Économie consolidée APU 2036 (M€) — moteur", None, "eco_consolidee_2036_meur"),
             ("Investissements (M€) — tableur", "invest_{}"),
             ("Cumul net 2027-2036 (M€) — moteur", None, "cumul_net_10ans_meur")]
    for i, line in enumerate(lines, 2):
        sy.cell(i, 1, line[0])
        for j, s in enumerate(SC, 2):
            sy.cell(i, j, f("=" + line[1].format(s)) if line[1] else round(float(synth[s][line[2]]), 1))
    k = len(lines) + 2
    sy.cell(k, 1, "Écart tableur / moteur sur les postes (%)")
    for j in range(2, 6):
        L = get_column_letter(j)
        sy.cell(k, j, f"=IF({L}3=0,0,({L}2-{L}3)/{L}3)").number_format = "0.0%"
    sy.cell(k + 1, 1, "Le tableur est un régime permanent sans calendrier : un petit écart avec le moteur (cohortes, montée en charge, mi-année) est attendu.")

    s0 = k + 3
    sy.cell(s0, 1, "Test des seuils (recalculé avec les hypothèses du classeur)").font = BOLD
    header(sy, s0 + 1, ["Seuil (postes)", "Atteint en intermédiaire ?", "Atteint en radical ?", "Capacité de départs sur l'horizon", "Total des viviers", "Verdict tableur", "Verdict de la note 12 (hypothèses centrales)"])
    for i, srow in enumerate(seuils_csv, s0 + 2):
        sv = int(srow["seuil"])
        sy.cell(i, 1, sv)
        sy.cell(i, 2, f(f'=IF(postes_intermediaire>=A{i},"oui","non")'))
        sy.cell(i, 3, f(f'=IF(postes_radical>=A{i},"oui","non")'))
        sy.cell(i, 4, f"=Viviers!I{nv + 1}")
        sy.cell(i, 5, f"=Viviers!C{nv + 1}")
        sy.cell(i, 6, f'=IF(B{i}="oui","conditionnel",IF(C{i}="oui","conditionnel",IF(D{i}>=A{i},"non démontré","incompatible")))')
        sy.cell(i, 7, srow["verdict"])
    widths(sy, [62, 18, 18, 18, 18, 20, 30])

    # --- Sensibilités (résultats du moteur)
    se = wb.create_sheet("Sensibilites")
    scols = ["variante", "sens", "libelle", "scenario", "postes_2036", "net_2036", "cumul_net_10", "cumul_consolide_10", "invest"]
    header(se, 1, scols + ["(résultats de scripts/12_sensibilites.py, moteur annuel)"])
    for i, row in enumerate(sens, 2):
        for j, kk in enumerate(scols, 1):
            val = row[kk]
            se.cell(i, j, float(val) if kk in scols[4:] else val)
    widths(se, [36, 6, 36, 14, 12, 12, 12, 14, 12])

    for sh in wb.worksheets:
        sh.freeze_panes = "A2" if sh.title != "Lisez-moi" else None
    out = ROOT / "modele/modele_financier.xlsx"
    wb.save(out)
    print(f"{out.relative_to(ROOT)} : {len(wb.worksheets)} feuilles, {len(reformes)} réformes")


if __name__ == "__main__":
    main()
